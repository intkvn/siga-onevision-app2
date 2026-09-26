import io
import json
import math
import os
import shutil
import tempfile
from datetime import datetime
from types import SimpleNamespace
from urllib.parse import quote_plus, urlencode

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import case, func, or_
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session, joinedload, selectinload

from app.auth import requiere_administrador
from app.config import ANIO_INVENTARIO
from app.database import get_db
from app.models import (
    BienInventarioImpresion,
    CargaInventarioImpresion,
    InventarioImpresion,
    ItemLoteImpresionInventario,
    ItemSolicitudImpresionInventario,
    LoteImpresionInventario,
    SolicitudImpresionInventario,
)
from app.routers.impresion import _obtener_perfil_impresion
from app.services.excel_inventario_impresion import importar_reporte_inventario
from app.services.pagination import paginas_visibles, rango_registros
from app.services.pdf_etiquetas import clasificar_bienes_impresion, generar_pdf_etiquetas
from app.services.solicitudes_impresion import (
    ESTADO_SINCRONIZACION,
    actualizar_estado_solicitud,
)


router = APIRouter()
templates = Jinja2Templates(directory="app/templates")

FILAS_POR_PAGINA = 50
MAX_BIENES_POR_LOTE = 1000
SIN_DATO = "__SIN_DATO__"


def _filtrar_valor(consulta, columna, valor):
    if not valor:
        return consulta
    if valor == SIN_DATO:
        return consulta.filter(or_(columna.is_(None), columna == ""))
    return consulta.filter(columna == valor)


def _aplicar_filtros(
    consulta,
    inventario_id: int,
    red: str = "",
    establecimiento: str = "",
    area: str = "",
    estado: str = "",
    tipo_bien: str = "",
    q: str = "",
):
    consulta = consulta.filter(
        BienInventarioImpresion.inventario_id == inventario_id,
        BienInventarioImpresion.activo == 1,
    )
    consulta = _filtrar_valor(consulta, BienInventarioImpresion.red, red)
    consulta = _filtrar_valor(
        consulta, BienInventarioImpresion.establecimiento, establecimiento
    )
    consulta = _filtrar_valor(consulta, BienInventarioImpresion.area, area)
    if estado:
        consulta = consulta.filter(BienInventarioImpresion.estado_impresion == estado)
    if tipo_bien:
        consulta = consulta.filter(BienInventarioImpresion.tipo_bien == tipo_bien)
    q = q.strip()
    if q:
        condiciones = [
            BienInventarioImpresion.codigo_patrimonial.ilike(f"{q}%"),
            BienInventarioImpresion.codigo_qr.ilike(f"{q}%"),
        ]
        if len(q) >= 3:
            condiciones.append(BienInventarioImpresion.descripcion.ilike(f"%{q}%"))
        consulta = consulta.filter(or_(
            *condiciones,
        ))
    return consulta


def _opciones_distintas(
    db, columna, inventario_id, q: str = "", limite: int = 50, **filtros
):
    consulta = db.query(columna).filter(
        BienInventarioImpresion.inventario_id == inventario_id,
        BienInventarioImpresion.activo == 1,
    )
    for nombre, valor in filtros.items():
        if nombre == "red":
            consulta = _filtrar_valor(consulta, BienInventarioImpresion.red, valor)
        elif nombre == "establecimiento":
            consulta = _filtrar_valor(
                consulta, BienInventarioImpresion.establecimiento, valor
            )
    q = q.strip()
    if q:
        consulta = consulta.filter(columna.ilike(f"%{q}%"))
    valores = [
        fila[0]
        for fila in consulta.distinct().order_by(columna).limit(limite).all()
    ]
    opciones = []
    if not q and any(valor is None or valor == "" for valor in valores):
        opciones.append({"value": SIN_DATO, "label": "Sin dato"})
    opciones.extend(
        {"value": valor, "label": valor}
        for valor in valores if valor is not None and valor != ""
    )
    return opciones


def _estado_bien(bien) -> str:
    return bien.estado_impresion


def _ultima_impresion(bien):
    return bien.impreso_en


def _adaptar_bien_pdf(bien):
    return SimpleNamespace(
        id=bien.id,
        codigo_patrimonial=bien.codigo_patrimonial,
        codigo_qr=bien.codigo_qr,
        ruta_qr=bien.ruta_qr,
        descripcion=bien.descripcion,
        centro_costo=SimpleNamespace(nombre_depend=bien.establecimiento or ""),
    )


def _marcar_pdf_generado(db: Session, lote: LoteImpresionInventario):
    ahora = datetime.utcnow()
    if lote.pdf_generado_en is None:
        lote.pdf_generado_en = ahora
    lote.estado = "Sticker generado"
    ids_bienes = [item.bien_id for item in lote.items]
    if ids_bienes:
        db.query(BienInventarioImpresion).filter(
            BienInventarioImpresion.id.in_(ids_bienes),
            BienInventarioImpresion.estado_impresion != "Impreso",
        ).update({
            BienInventarioImpresion.estado_impresion: "Sticker generado",
            BienInventarioImpresion.sticker_generado_en: ahora,
        }, synchronize_session=False)
    db.commit()


def _confirmar_items_impresos(db: Session, lote, consulta) -> int:
    pendientes_items = consulta.filter(
        ItemLoteImpresionInventario.impreso_en.is_(None)
    ).all()
    ids_bienes = [item.bien_id for item in pendientes_items]
    ids_solicitud = [
        item.solicitud_item_id for item in pendientes_items
        if item.solicitud_item_id is not None
    ]
    ahora = datetime.utcnow()
    for item in pendientes_items:
        item.impreso_en = ahora
    if ids_bienes:
        db.query(BienInventarioImpresion).filter(
            BienInventarioImpresion.id.in_(ids_bienes)
        ).update({
            BienInventarioImpresion.estado_impresion: "Impreso",
            BienInventarioImpresion.impreso_en: ahora,
        }, synchronize_session=False)
    solicitudes_afectadas = set()
    if ids_solicitud:
        items_solicitud = db.query(ItemSolicitudImpresionInventario).filter(
            ItemSolicitudImpresionInventario.id.in_(ids_solicitud)
        ).all()
        for item in items_solicitud:
            item.estado = "Listo para recojo"
            item.listo_recojo_en = ahora
            solicitudes_afectadas.add(item.solicitud_id)
        db.flush()
        for solicitud_id in solicitudes_afectadas:
            actualizar_estado_solicitud(
                db.get(SolicitudImpresionInventario, solicitud_id)
            )
    db.flush()
    pendientes = db.query(ItemLoteImpresionInventario).filter(
        ItemLoteImpresionInventario.lote_id == lote.id,
        ItemLoteImpresionInventario.impreso_en.is_(None),
    ).count()
    lote.estado = "Impreso" if pendientes == 0 else "Impreso parcial"
    db.commit()
    return len(pendientes_items)


def _resumen_establecimientos(consulta):
    estado = BienInventarioImpresion.estado_impresion
    filas = (
        consulta.with_entities(
            BienInventarioImpresion.red,
            BienInventarioImpresion.establecimiento,
            func.count(BienInventarioImpresion.id),
            func.sum(case((estado == "Pendiente", 1), else_=0)),
            func.sum(case((estado == "Sticker generado", 1), else_=0)),
            func.sum(case((estado == "Impreso", 1), else_=0)),
            func.sum(case((estado == "Bloqueado", 1), else_=0)),
        )
        .group_by(
            BienInventarioImpresion.red,
            BienInventarioImpresion.establecimiento,
        )
        .order_by(
            BienInventarioImpresion.red,
            BienInventarioImpresion.establecimiento,
        )
        .all()
    )
    return [
        {
            "red": red or "Sin RED",
            "establecimiento": establecimiento or "Sin establecimiento",
            "total": total,
            "pendientes": pendientes or 0,
            "generados": generados or 0,
            "impresos": impresos or 0,
            "bloqueados": bloqueados or 0,
            "porcentaje": round((impresos or 0) * 100 / total, 1) if total else 0,
            "completo": bool(total and impresos == total),
        }
        for red, establecimiento, total, pendientes, generados, impresos, bloqueados
        in filas
    ]


@router.get("/control-impresion/opciones")
def opciones_filtro_impresion(
    inventario_id: int,
    campo: str,
    q: str = "",
    red: str = "",
    establecimiento: str = "",
    db: Session = Depends(get_db),
    _=Depends(requiere_administrador),
):
    columnas = {
        "red": BienInventarioImpresion.red,
        "establecimiento": BienInventarioImpresion.establecimiento,
        "area": BienInventarioImpresion.area,
    }
    columna = columnas.get(campo)
    if columna is None:
        raise HTTPException(status_code=400, detail="Filtro no válido.")
    filtros = {}
    if campo in ("establecimiento", "area"):
        filtros["red"] = red
    if campo == "area":
        filtros["establecimiento"] = establecimiento
    return {
        "options": _opciones_distintas(
            db, columna, inventario_id, q=q, limite=50, **filtros
        )
    }


@router.get("/control-impresion", response_class=HTMLResponse)
def control_impresion(
    request: Request,
    inventario_id: int | None = None,
    red: str = "",
    establecimiento: str = "",
    area: str = "",
    estado: str = "",
    tipo_bien: str = "",
    q: str = "",
    pagina: int = 1,
    db: Session = Depends(get_db),
    _=Depends(requiere_administrador),
):
    inventarios = db.query(InventarioImpresion).order_by(
        InventarioImpresion.anio.desc(), InventarioImpresion.id.desc()
    ).all()
    inventario = None
    if inventario_id:
        inventario = db.get(InventarioImpresion, inventario_id)
    if inventario is None and inventarios:
        inventario = inventarios[0]

    contexto = {
        "request": request,
        "inventarios": inventarios,
        "inventario": inventario,
        "anio_predeterminado": ANIO_INVENTARIO,
        "max_bienes_lote": MAX_BIENES_POR_LOTE,
        "estados": ("Pendiente", "Sticker generado", "Impreso", "Bloqueado"),
        "filtros": {
            "red": red,
            "establecimiento": establecimiento,
            "area": area,
            "estado": estado,
            "tipo_bien": tipo_bien,
            "q": q,
        },
        "bienes": [],
        "resumen": {nombre: 0 for nombre in (
            "total", "Pendiente", "Sticker generado", "Impreso", "Bloqueado"
        )},
        "resumen_establecimientos": [],
        "lotes": [],
        "cargas_reportes": [],
        "tipos": {"Activo fijo": 0, "Sobrante": 0},
        "solicitudes_pendientes": 0,
        "solicitudes_sincronizacion": 0,
        "pagina": 1,
        "total_paginas": 1,
        "paginas": [1],
        "inicio": 0,
        "fin": 0,
        "total": 0,
        "url_base": "/control-impresion?",
    }
    if inventario is None:
        return templates.TemplateResponse("control_impresion.html", contexto)

    base = _aplicar_filtros(
        db.query(BienInventarioImpresion), inventario.id,
        red=red, establecimiento=establecimiento, area=area, estado=estado,
        tipo_bien=tipo_bien, q=q,
    )
    total = base.count()
    total_paginas = max(1, math.ceil(total / FILAS_POR_PAGINA))
    pagina = max(1, min(pagina, total_paginas))
    bienes = (
        base.options(
            joinedload(BienInventarioImpresion.bien_alta),
        )
        .order_by(
            BienInventarioImpresion.establecimiento,
            BienInventarioImpresion.area,
            BienInventarioImpresion.codigo_patrimonial,
        )
        .offset((pagina - 1) * FILAS_POR_PAGINA)
        .limit(FILAS_POR_PAGINA)
        .all()
    )
    filas_bienes = [
        {
            "bien": bien,
            "estado": _estado_bien(bien),
            "ultima_impresion": _ultima_impresion(bien),
        }
        for bien in bienes
    ]

    conteos = {nombre: 0 for nombre in (
        "Pendiente", "Sticker generado", "Impreso", "Bloqueado"
    )}
    for nombre, cantidad in (
        base.with_entities(
            BienInventarioImpresion.estado_impresion,
            func.count(BienInventarioImpresion.id),
        )
        .group_by(BienInventarioImpresion.estado_impresion)
        .all()
    ):
        conteos[nombre] = cantidad

    parametros = {
        "inventario_id": inventario.id,
        "red": red,
        "establecimiento": establecimiento,
        "area": area,
        "estado": estado,
        "tipo_bien": tipo_bien,
        "q": q,
    }
    url_base = "/control-impresion?" + urlencode(
        {clave: valor for clave, valor in parametros.items() if valor not in (None, "")}
    )
    inicio, fin = rango_registros(pagina, FILAS_POR_PAGINA, total)
    contexto.update({
        "bienes": filas_bienes,
        "resumen": {"total": total, **conteos},
        "tipos": {
            **{"Activo fijo": 0, "Sobrante": 0},
            **dict(
                base.with_entities(
                    BienInventarioImpresion.tipo_bien,
                    func.count(BienInventarioImpresion.id),
                ).group_by(BienInventarioImpresion.tipo_bien).all()
            ),
        },
        "solicitudes_pendientes": db.query(
            ItemSolicitudImpresionInventario
        ).filter(
            ItemSolicitudImpresionInventario.estado == "Pendiente"
        ).count(),
        "solicitudes_sincronizacion": db.query(
            ItemSolicitudImpresionInventario
        ).filter(
            ItemSolicitudImpresionInventario.estado == ESTADO_SINCRONIZACION
        ).count(),
        "resumen_establecimientos": _resumen_establecimientos(base),
        "lotes": (
            db.query(LoteImpresionInventario)
            .filter(LoteImpresionInventario.inventario_id == inventario.id)
            .order_by(LoteImpresionInventario.id.desc())
            .limit(20)
            .all()
        ),
        "cargas_reportes": (
            db.query(CargaInventarioImpresion)
            .filter(CargaInventarioImpresion.inventario_id == inventario.id)
            .order_by(
                CargaInventarioImpresion.creado_en.desc(),
                CargaInventarioImpresion.id.desc(),
            )
            .limit(50)
            .all()
        ),
        "pagina": pagina,
        "total_paginas": total_paginas,
        "paginas": paginas_visibles(pagina, total_paginas),
        "inicio": inicio,
        "fin": fin,
        "total": total,
        "url_base": url_base,
    })
    return templates.TemplateResponse("control_impresion.html", contexto)


@router.post("/control-impresion/importar")
def importar_inventario(
    archivo: UploadFile = File(...),
    anio: str = Form(ANIO_INVENTARIO),
    db: Session = Depends(get_db),
    _=Depends(requiere_administrador),
):
    ruta_temporal = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
            ruta_temporal = tmp.name
            shutil.copyfileobj(archivo.file, tmp, length=1024 * 1024)
        resultado = importar_reporte_inventario(
            db, ruta_temporal, archivo.filename or "reporte.xlsx", anio=anio,
        )
    except (ValueError, OSError) as exc:
        db.rollback()
        return RedirectResponse(
            url=f"/control-impresion?error={quote_plus(str(exc))}", status_code=303
        )
    except SQLAlchemyError:
        db.rollback()
        return RedirectResponse(
            url=(
                "/control-impresion?error="
                + quote_plus(
                    "No se pudo completar la carga por un conflicto entre "
                    "identificadores. Revisa los códigos QR del archivo."
                )
            ),
            status_code=303,
        )
    finally:
        if ruta_temporal and os.path.exists(ruta_temporal):
            os.remove(ruta_temporal)

    mensaje = (
        f"Importación acumulativa completada: {resultado['total']} filas "
        f"procesadas, {resultado['nuevos']} nuevas, "
        f"{resultado['actualizados']} actualizadas y "
        f"{resultado['sin_cambios']} sin cambios. Universo actual: "
        f"{resultado['total_universo']} bienes, "
        f"{resultado['activos_fijos_universo']} activos fijos y "
        f"{resultado['sobrantes_universo']} sobrantes."
    )
    if resultado["solicitudes_vinculadas"]:
        mensaje += (
            f" {resultado['solicitudes_vinculadas']} QR solicitado(s) quedaron "
            "listos para preparar su impresión."
        )
    return RedirectResponse(
        url=(
            f"/control-impresion?inventario_id={resultado['inventario_id']}"
            f"&info={quote_plus(mensaje)}"
        ),
        status_code=303,
    )


@router.post("/control-impresion/lotes")
def crear_lote_impresion(
    inventario_id: int = Form(...),
    alcance: str = Form(...),
    bien_ids: list[int] = Form(default=[]),
    red: str = Form(""),
    establecimiento: str = Form(""),
    area: str = Form(""),
    estado: str = Form(""),
    tipo_bien: str = Form(""),
    q: str = Form(""),
    db: Session = Depends(get_db),
    _=Depends(requiere_administrador),
):
    inventario = db.get(InventarioImpresion, inventario_id)
    if inventario is None:
        raise HTTPException(status_code=404, detail="Inventario no encontrado.")
    consulta = _aplicar_filtros(
        db.query(BienInventarioImpresion), inventario_id,
        red=red, establecimiento=establecimiento, area=area, estado=estado,
        tipo_bien=tipo_bien, q=q,
    )
    if alcance == "seleccionados":
        if not bien_ids:
            return RedirectResponse(
                url=f"/control-impresion?inventario_id={inventario_id}&error="
                    + quote_plus("Selecciona al menos un bien."),
                status_code=303,
            )
        consulta = consulta.filter(BienInventarioImpresion.id.in_(set(bien_ids)))
    elif alcance != "filtrados":
        raise HTTPException(status_code=400, detail="Alcance no válido.")

    bienes = consulta.order_by(BienInventarioImpresion.id).all()
    if len(bienes) > MAX_BIENES_POR_LOTE:
        mensaje = (
            f"La selección tiene {len(bienes)} bienes. El máximo por lote es "
            f"{MAX_BIENES_POR_LOTE}; aplica más filtros."
        )
        return RedirectResponse(
            url=f"/control-impresion?inventario_id={inventario_id}&error="
                + quote_plus(mensaje),
            status_code=303,
        )

    adaptados = [_adaptar_bien_pdf(bien) for bien in bienes]
    imprimibles, excluidos = clasificar_bienes_impresion(adaptados)
    ids_imprimibles = {bien.id for bien in imprimibles}
    bienes = [bien for bien in bienes if bien.id in ids_imprimibles]
    if not bienes:
        return RedirectResponse(
            url=f"/control-impresion?inventario_id={inventario_id}&error="
                + quote_plus("No hay bienes imprimibles en la selección."),
            status_code=303,
        )

    filtros = {
        "red": red or None,
        "establecimiento": establecimiento or None,
        "area": area or None,
        "estado": estado or None,
        "tipo_bien": tipo_bien or None,
        "busqueda": q or None,
        "excluidos": len(excluidos),
    }
    lote = LoteImpresionInventario(
        inventario_id=inventario_id,
        estado="Preparado",
        filtros=json.dumps(filtros, ensure_ascii=False),
        total_bienes=len(bienes),
    )
    db.add(lote)
    db.flush()
    db.bulk_insert_mappings(ItemLoteImpresionInventario, [
        {"lote_id": lote.id, "bien_id": bien.id} for bien in bienes
    ])
    db.commit()
    info = f"Lote preparado con {len(bienes)} bienes."
    if excluidos:
        info += f" Se excluyeron {len(excluidos)} bienes no imprimibles."
    return RedirectResponse(
        url=f"/control-impresion/lotes/{lote.id}?info={quote_plus(info)}",
        status_code=303,
    )


@router.get("/control-impresion/solicitudes", response_class=HTMLResponse)
def solicitudes_inventariadores(
    request: Request,
    db: Session = Depends(get_db),
    _=Depends(requiere_administrador),
):
    solicitudes = (
        db.query(SolicitudImpresionInventario)
        .options(
            joinedload(SolicitudImpresionInventario.usuario),
            selectinload(SolicitudImpresionInventario.items)
            .joinedload(ItemSolicitudImpresionInventario.bien),
        )
        .order_by(SolicitudImpresionInventario.id.desc())
        .limit(200)
        .all()
    )
    pendientes = [
        item for solicitud in solicitudes for item in solicitud.items
        if item.estado == "Pendiente"
    ]
    esperando = [
        item for solicitud in solicitudes for item in solicitud.items
        if item.estado == ESTADO_SINCRONIZACION
    ]
    return templates.TemplateResponse(
        "control_impresion_solicitudes.html",
        {
            "request": request,
            "solicitudes": solicitudes,
            "pendientes": pendientes,
            "esperando": esperando,
            "max_bienes_lote": MAX_BIENES_POR_LOTE,
        },
    )


@router.post("/control-impresion/solicitudes/lote")
def crear_lote_desde_solicitudes(
    item_ids: list[int] = Form(default=[]),
    db: Session = Depends(get_db),
    _=Depends(requiere_administrador),
):
    if not item_ids:
        return RedirectResponse(
            "/control-impresion/solicitudes?error="
            + quote_plus("Selecciona al menos un QR pendiente."),
            status_code=303,
        )
    items = (
        db.query(ItemSolicitudImpresionInventario)
        .options(
            joinedload(ItemSolicitudImpresionInventario.bien),
            joinedload(ItemSolicitudImpresionInventario.solicitud),
        )
        .filter(
            ItemSolicitudImpresionInventario.id.in_(set(item_ids)),
            ItemSolicitudImpresionInventario.estado == "Pendiente",
        )
        .order_by(ItemSolicitudImpresionInventario.id)
        .all()
    )
    if not items:
        return RedirectResponse(
            "/control-impresion/solicitudes?error="
            + quote_plus("Los QR seleccionados ya no están pendientes."),
            status_code=303,
        )
    inventarios = {item.solicitud.inventario_id for item in items}
    if len(inventarios) != 1:
        return RedirectResponse(
            "/control-impresion/solicitudes?error="
            + quote_plus("Selecciona solicitudes de un solo inventario."),
            status_code=303,
        )
    if len(items) > MAX_BIENES_POR_LOTE:
        return RedirectResponse(
            "/control-impresion/solicitudes?error=" + quote_plus(
                f"El máximo por lote es {MAX_BIENES_POR_LOTE} QR."
            ), status_code=303,
        )

    seleccionados = []
    bienes_vistos = set()
    solicitudes_afectadas = set()
    for item in items:
        solicitudes_afectadas.add(item.solicitud_id)
        if item.bien is None:
            item.estado = "Observado"
            item.motivo_observacion = "El QR ya no tiene un bien relacionado."
            continue
        if item.bien_id in bienes_vistos:
            item.estado = "Observado"
            item.motivo_observacion = "El mismo bien fue incluido más de una vez."
            continue
        razon = clasificar_bienes_impresion([_adaptar_bien_pdf(item.bien)])[1]
        if razon:
            item.estado = "Observado"
            item.motivo_observacion = razon[0]["razon"]
            continue
        bienes_vistos.add(item.bien_id)
        seleccionados.append(item)

    if not seleccionados:
        for solicitud_id in solicitudes_afectadas:
            actualizar_estado_solicitud(db.get(SolicitudImpresionInventario, solicitud_id))
        db.commit()
        return RedirectResponse(
            "/control-impresion/solicitudes?error="
            + quote_plus("Ninguno de los QR seleccionados puede imprimirse."),
            status_code=303,
        )

    inventario_id = seleccionados[0].solicitud.inventario_id
    lote = LoteImpresionInventario(
        inventario_id=inventario_id,
        estado="Preparado",
        filtros=json.dumps({"origen": "Solicitudes de inventariadores"}),
        total_bienes=len(seleccionados),
    )
    db.add(lote)
    db.flush()
    for item in seleccionados:
        db.add(ItemLoteImpresionInventario(
            lote_id=lote.id,
            bien_id=item.bien_id,
            solicitud_item_id=item.id,
        ))
        item.estado = "En lote"
    db.flush()
    for solicitud_id in solicitudes_afectadas:
        actualizar_estado_solicitud(db.get(SolicitudImpresionInventario, solicitud_id))
    db.commit()
    return RedirectResponse(
        f"/control-impresion/lotes/{lote.id}?info=" + quote_plus(
            f"Lote preparado con {len(seleccionados)} QR solicitados."
        ), status_code=303,
    )


@router.get("/control-impresion/lotes/{lote_id}", response_class=HTMLResponse)
def detalle_lote_impresion(
    lote_id: int,
    request: Request,
    db: Session = Depends(get_db),
    _=Depends(requiere_administrador),
):
    lote = (
        db.query(LoteImpresionInventario)
        .options(
            joinedload(LoteImpresionInventario.inventario),
            selectinload(LoteImpresionInventario.items)
            .joinedload(ItemLoteImpresionInventario.bien)
            .joinedload(BienInventarioImpresion.bien_alta),
        )
        .filter(LoteImpresionInventario.id == lote_id)
        .first()
    )
    if lote is None:
        raise HTTPException(status_code=404, detail="Lote no encontrado.")
    impresos = sum(1 for item in lote.items if item.impreso_en)
    return templates.TemplateResponse(
        "control_impresion_lote.html",
        {
            "request": request,
            "lote": lote,
            "impresos": impresos,
            "pendientes": len(lote.items) - impresos,
        },
    )


@router.get("/control-impresion/lotes/{lote_id}/pdf")
def pdf_lote_impresion(
    lote_id: int,
    db: Session = Depends(get_db),
    _=Depends(requiere_administrador),
):
    lote = (
        db.query(LoteImpresionInventario)
        .options(
            selectinload(LoteImpresionInventario.items)
            .joinedload(ItemLoteImpresionInventario.bien)
        )
        .filter(LoteImpresionInventario.id == lote_id)
        .first()
    )
    if lote is None:
        raise HTTPException(status_code=404, detail="Lote no encontrado.")
    bienes = [_adaptar_bien_pdf(item.bien) for item in lote.items]
    perfil = _obtener_perfil_impresion(db)
    contenido = generar_pdf_etiquetas(bienes, perfil)
    _marcar_pdf_generado(db, lote)
    return Response(
        content=contenido,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'inline; filename="inventario_lote_{lote.id}.pdf"',
            "Cache-Control": "no-store",
        },
    )


def _crear_excel_lote(lote) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "BarTender"
    encabezados = [
        "Codigo Patrimonial", "Codigo QR", "Ruta QR", "Bien",
        "Establecimiento", "RED", "Area", "Marca", "Modelo", "Color",
        "Nro Serie",
    ]
    hoja.append(encabezados)
    for item in lote.items:
        bien = item.bien
        hoja.append([
            bien.codigo_patrimonial, bien.codigo_qr or "", bien.ruta_qr or "",
            bien.descripcion, bien.establecimiento or "", bien.red or "",
            bien.area or "", bien.marca or "", bien.modelo or "",
            bien.color or "", bien.nro_serie or "",
        ])

    control = libro.create_sheet("Control")
    control.append([
        "Item", "Codigo Patrimonial", "Codigo QR", "Bien", "RED",
        "Establecimiento", "Area", "Estado impresión", "Fecha impresión",
    ])
    for numero, item in enumerate(lote.items, start=1):
        bien = item.bien
        control.append([
            numero, bien.codigo_patrimonial, bien.codigo_qr or "", bien.descripcion,
            bien.red or "Sin RED", bien.establecimiento or "Sin establecimiento",
            bien.area or "Sin área", "Impreso" if item.impreso_en else "Pendiente",
            item.impreso_en,
        ])
    for hoja_actual in (hoja, control):
        for celda in hoja_actual[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F4E78")
            celda.alignment = Alignment(horizontal="center", vertical="center")
        hoja_actual.freeze_panes = "A2"
        hoja_actual.auto_filter.ref = hoja_actual.dimensions
        hoja_actual.sheet_view.showGridLines = False
    for columna, ancho in enumerate((22, 18, 55, 40, 32, 25, 30, 18, 18, 15, 20), 1):
        hoja.column_dimensions[hoja.cell(1, columna).column_letter].width = ancho
    for columna, ancho in enumerate((8, 22, 18, 40, 25, 32, 30, 18, 22), 1):
        control.column_dimensions[control.cell(1, columna).column_letter].width = ancho

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


@router.get("/control-impresion/lotes/{lote_id}/excel")
def excel_lote_impresion(
    lote_id: int,
    db: Session = Depends(get_db),
    _=Depends(requiere_administrador),
):
    lote = (
        db.query(LoteImpresionInventario)
        .options(
            selectinload(LoteImpresionInventario.items)
            .joinedload(ItemLoteImpresionInventario.bien)
        )
        .filter(LoteImpresionInventario.id == lote_id)
        .first()
    )
    if lote is None:
        raise HTTPException(status_code=404, detail="Lote no encontrado.")
    contenido = _crear_excel_lote(lote)
    return Response(
        content=contenido,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="inventario_lote_{lote.id}.xlsx"'
            )
        },
    )


@router.post("/control-impresion/lotes/{lote_id}/confirmar")
def confirmar_impresion_lote(
    lote_id: int,
    alcance: str = Form(...),
    item_ids: list[int] = Form(default=[]),
    db: Session = Depends(get_db),
    _=Depends(requiere_administrador),
):
    lote = db.get(LoteImpresionInventario, lote_id)
    if lote is None:
        raise HTTPException(status_code=404, detail="Lote no encontrado.")
    if lote.pdf_generado_en is None:
        return RedirectResponse(
            url=f"/control-impresion/lotes/{lote_id}?error="
                + quote_plus("Genera primero el PDF del lote."),
            status_code=303,
        )
    consulta = db.query(ItemLoteImpresionInventario).filter(
        ItemLoteImpresionInventario.lote_id == lote_id
    )
    if alcance == "seleccionados":
        if not item_ids:
            return RedirectResponse(
                url=f"/control-impresion/lotes/{lote_id}?error="
                    + quote_plus("Selecciona al menos un bien."),
                status_code=303,
            )
        consulta = consulta.filter(ItemLoteImpresionInventario.id.in_(set(item_ids)))
    elif alcance != "todo":
        raise HTTPException(status_code=400, detail="Alcance no válido.")

    actualizados = _confirmar_items_impresos(db, lote, consulta)
    return RedirectResponse(
        url=(
            f"/control-impresion/lotes/{lote_id}?info="
            + quote_plus(f"Se confirmaron {actualizados} bienes como impresos.")
        ),
        status_code=303,
    )
