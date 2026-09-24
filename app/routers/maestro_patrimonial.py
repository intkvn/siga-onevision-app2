import io
import json
import math
import os
import re
import shutil
import tempfile
import threading
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from urllib.parse import urlencode

from fastapi import (
    APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, Request,
    UploadFile,
)
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, Response
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_
from sqlalchemy.orm import Session, selectinload

from app.auth import requiere_login
from app.database import SessionLocal, get_db
from app.models import (
    BienAlta,
    BienCargaPatrimonial,
    BienInventarioImpresion,
    BienPatrimonial,
    CargaPatrimonial,
    ConflictoBienPatrimonial,
    CorreccionBienPatrimonial,
    ExportacionPatrimonial,
    Pecosa,
    Persona,
    VersionBienPatrimonial,
)
from app.services.maestro_patrimonial import (
    CAMPOS,
    CAMPOS_EDITABLES,
    CAMPOS_OPCIONALES,
    ENCABEZADOS_SIGA,
    ETIQUETAS_CAMPOS,
    confirmar_carga_patrimonial,
    editar_bien_patrimonial,
    normalizar_codigo_qr,
    preparar_carga,
    resolver_conflicto,
    serializar_valor,
    validar_carga_patrimonial_desde_bd,
)
from app.services.pagination import paginas_visibles, rango_registros
from app.services.pdf_ficha_patrimonial import generar_ficha_activo_pdf


router = APIRouter()
templates = Jinja2Templates(directory="app/templates")
FILAS_POR_PAGINA = 50
FILAS_CARGA_POR_PAGINA = 100
TIPOS_COLUMNAS_EXPORTACION = {
    posicion: tipo for posicion, tipo in CAMPOS.values()
}

FILTROS_TEXTO = {
    "dependencia": BienPatrimonial.nombre_dependencia,
    "usuario": BienPatrimonial.usuario,
    "ubicacion": BienPatrimonial.ubicacion_fisica,
    "modelo": BienPatrimonial.modelo,
    "numero_orden": BienPatrimonial.numero_orden,
    "numero_documento": BienPatrimonial.numero_documento,
    "marca": BienPatrimonial.marca,
    "estado_conservacion": BienPatrimonial.estado_conservacion,
    "numero_serie": BienPatrimonial.numero_serie,
    "color": BienPatrimonial.color,
}
FILTROS_FECHA = {
    "fecha_compra": BienPatrimonial.fecha_compra,
    "fecha_alta": BienPatrimonial.fecha_alta,
    "fecha_nea": BienPatrimonial.fecha_nea,
}
FILTROS_OPCIONES = {
    "dependencia": BienPatrimonial.nombre_dependencia,
    "usuario": BienPatrimonial.usuario,
    "ubicacion": BienPatrimonial.ubicacion_fisica,
    "modelo": BienPatrimonial.modelo,
    "marca": BienPatrimonial.marca,
    "estado_conservacion": BienPatrimonial.estado_conservacion,
    "color": BienPatrimonial.color,
}


def _usuario(request: Request) -> str:
    return request.session.get("usuario") or "admin"


def _fecha_filtro(valor: str):
    try:
        return date.fromisoformat(valor) if valor else None
    except ValueError:
        return None


def _filtros_desde_parametros(**valores) -> dict:
    return {clave: (valor.strip() if isinstance(valor, str) else valor) for clave, valor in valores.items()}


def _aplicar_filtros(consulta, filtros: dict):
    codigo_patrimonial = filtros.get("codigo_patrimonial", "").strip()
    if codigo_patrimonial:
        consulta = consulta.filter(
            func.lower(BienPatrimonial.codigo_patrimonial).like(
                f"{codigo_patrimonial.lower()}%"
            )
        )
    codigo_qr = filtros.get("codigo_qr", "").strip()
    if codigo_qr:
        codigo_qr = normalizar_codigo_qr(codigo_qr) or ""
        consulta = consulta.filter(
            func.lower(BienPatrimonial.codigo_qr) == codigo_qr.lower()
        )
    descripcion = filtros.get("descripcion", "").strip()
    if descripcion:
        consulta = consulta.filter(BienPatrimonial.descripcion.ilike(f"%{descripcion}%"))

    # Conserva compatibilidad con enlaces anteriores que usaban una sola búsqueda.
    q = filtros.get("q", "").strip()
    if q:
        condiciones = [
            func.lower(BienPatrimonial.codigo_patrimonial).like(f"{q.lower()}%"),
            func.lower(BienPatrimonial.codigo_qr).like(f"{q.lower()}%"),
            BienPatrimonial.descripcion.ilike(f"%{q}%"),
        ]
        consulta = consulta.filter(or_(*condiciones))
    for nombre, columna in FILTROS_TEXTO.items():
        valor = filtros.get(nombre, "")
        if valor:
            consulta = consulta.filter(columna.ilike(f"%{valor}%"))
    for nombre, columna in FILTROS_FECHA.items():
        desde = _fecha_filtro(filtros.get(f"{nombre}_desde", ""))
        hasta = _fecha_filtro(filtros.get(f"{nombre}_hasta", ""))
        if desde:
            consulta = consulta.filter(columna >= desde)
        if hasta:
            consulta = consulta.filter(columna <= hasta)
    return consulta


def _parametros_filtros(request: Request) -> dict:
    nombres = [
        "codigo_patrimonial", "codigo_qr", "descripcion", "q", *FILTROS_TEXTO,
    ]
    nombres.extend(
        nombre for campo in FILTROS_FECHA for nombre in (f"{campo}_desde", f"{campo}_hasta")
    )
    return {nombre: request.query_params.get(nombre, "").strip() for nombre in nombres}


def _datos_firmante(db: Session, texto_firmante: str | None) -> tuple[str, str]:
    """Separa el texto histórico del firmante y completa el DNI desde Personas."""
    texto = (texto_firmante or "").strip()
    if not texto:
        return "", ""

    coincidencia = re.fullmatch(r"(.+?)\s*\(([^()]*)\)\s*", texto)
    if coincidencia:
        nombre = coincidencia.group(1).strip()
        dni = coincidencia.group(2).strip()
        return nombre, dni

    persona = db.query(Persona).filter(or_(
        Persona.nombre_completo.ilike(texto),
        Persona.dni == texto,
    )).first()
    if persona:
        return persona.nombre_completo, persona.dni
    return texto, ""


def _texto_identificador(valor) -> str:
    return str(valor or "").strip()


def _fecha_hora_evento(valor) -> datetime | None:
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime.combine(valor, datetime.min.time())
    return None


def _resumen_calidad_datos(db: Session) -> dict:
    maestros = db.query(
        BienPatrimonial.id,
        BienPatrimonial.codigo_patrimonial,
        BienPatrimonial.codigo_qr,
    ).all()
    altas = db.query(BienAlta.codigo_patrimonial, BienAlta.codigo_qr).all()
    impresiones = db.query(
        BienInventarioImpresion.codigo_patrimonial,
        BienInventarioImpresion.codigo_qr,
    ).all()

    id_por_codigo = {fila.codigo_patrimonial: fila.id for fila in maestros}
    qr_maestro = defaultdict(list)
    sin_qr = []
    por_codigo = defaultdict(lambda: defaultdict(set))
    qr_a_codigos = defaultdict(set)

    def registrar(origen: str, codigo, qr, bien_id=None):
        codigo_texto = _texto_identificador(codigo)
        qr_texto = normalizar_codigo_qr(qr) or ""
        if not codigo_texto:
            return
        if qr_texto:
            por_codigo[codigo_texto][origen].add(qr_texto)
            qr_a_codigos[qr_texto].add(codigo_texto)
        if origen == "Maestro":
            if qr_texto:
                qr_maestro[qr_texto].append({
                    "codigo": codigo_texto, "bien_id": bien_id,
                })
            else:
                sin_qr.append({"codigo": codigo_texto, "bien_id": bien_id})

    for fila in maestros:
        registrar("Maestro", fila.codigo_patrimonial, fila.codigo_qr, fila.id)
    for codigo, qr in altas:
        registrar("Altas", codigo, qr)
    for codigo, qr in impresiones:
        registrar("Control Impresión", codigo, qr)

    duplicados_maestro = [
        {"qr": qr, "bienes": bienes}
        for qr, bienes in sorted(qr_maestro.items()) if len(bienes) > 1
    ]
    discrepancias = []
    for codigo, origenes in sorted(por_codigo.items()):
        valores = set().union(*origenes.values()) if origenes else set()
        if len(valores) > 1:
            discrepancias.append({
                "codigo": codigo,
                "bien_id": id_por_codigo.get(codigo),
                "origenes": [
                    {"origen": origen, "qr": ", ".join(sorted(qrs))}
                    for origen, qrs in sorted(origenes.items())
                ],
            })
    qr_otros_codigos = [
        {"qr": qr, "codigos": sorted(codigos)}
        for qr, codigos in sorted(qr_a_codigos.items()) if len(codigos) > 1
    ]

    opcionales = []
    for campo in CAMPOS_OPCIONALES:
        columna = getattr(BienPatrimonial, campo)
        consulta = db.query(func.count(BienPatrimonial.id)).filter(columna.is_(None))
        if CAMPOS[campo][1] == "texto":
            consulta = db.query(func.count(BienPatrimonial.id)).filter(
                or_(columna.is_(None), columna == "")
            )
        opcionales.append({
            "campo": campo,
            "etiqueta": ETIQUETAS_CAMPOS[campo],
            "cantidad": consulta.scalar() or 0,
        })
    opcionales.sort(key=lambda item: (-item["cantidad"], item["etiqueta"]))

    return {
        "total_bienes": len(maestros),
        "sin_qr_total": len(sin_qr),
        "sin_qr": sin_qr[:100],
        "duplicados_total": len(duplicados_maestro),
        "duplicados": duplicados_maestro[:100],
        "discrepancias_total": len(discrepancias),
        "discrepancias": discrepancias[:200],
        "qr_otros_codigos_total": len(qr_otros_codigos),
        "qr_otros_codigos": qr_otros_codigos[:100],
        "opcionales": opcionales,
    }


@router.get("/maestro-patrimonial", response_class=HTMLResponse)
def maestro_patrimonial(
    request: Request,
    pagina: int = 1,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    filtros = _parametros_filtros(request)
    consulta = _aplicar_filtros(db.query(BienPatrimonial), filtros)
    total = consulta.count()
    total_paginas = max(1, math.ceil(total / FILAS_POR_PAGINA))
    pagina = max(1, min(pagina, total_paginas))
    bienes = consulta.order_by(
        BienPatrimonial.nombre_dependencia,
        BienPatrimonial.descripcion,
        BienPatrimonial.codigo_patrimonial,
    ).offset((pagina - 1) * FILAS_POR_PAGINA).limit(FILAS_POR_PAGINA).all()

    parametros = {k: v for k, v in filtros.items() if v}
    url_base = "/maestro-patrimonial?" + urlencode(parametros)
    inicio, fin = rango_registros(pagina, FILAS_POR_PAGINA, total)
    cargas = db.query(CargaPatrimonial).order_by(CargaPatrimonial.id.desc()).limit(12).all()
    exportaciones = db.query(ExportacionPatrimonial).order_by(
        ExportacionPatrimonial.id.desc()
    ).limit(8).all()
    pendientes = db.query(ConflictoBienPatrimonial).filter(
        ConflictoBienPatrimonial.estado == "Pendiente"
    ).count()
    ultima_carga = db.query(CargaPatrimonial).filter(
        CargaPatrimonial.estado == "Completada"
    ).order_by(CargaPatrimonial.id.desc()).first()
    return templates.TemplateResponse("maestro_patrimonial.html", {
        "request": request,
        "bienes": bienes,
        "cargas": cargas,
        "exportaciones": exportaciones,
        "ultima_carga": ultima_carga,
        "conflictos_pendientes": pendientes,
        "filtros": filtros,
        "pagina": pagina,
        "total_paginas": total_paginas,
        "paginas": paginas_visibles(pagina, total_paginas),
        "inicio": inicio,
        "fin": fin,
        "total": total,
        "url_base": url_base,
        "query_exportar": urlencode(parametros),
    })


@router.get("/maestro-patrimonial/opciones")
def opciones_maestro_patrimonial(
    campo: str,
    q: str = "",
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    columna = FILTROS_OPCIONES.get(campo)
    if columna is None:
        raise HTTPException(status_code=400, detail="Filtro no válido.")
    consulta = db.query(columna).filter(columna.is_not(None), columna != "")
    if q.strip():
        consulta = consulta.filter(columna.ilike(f"%{q.strip()}%"))
    valores = [fila[0] for fila in consulta.distinct().order_by(columna).limit(50).all()]
    return {"options": [{"value": valor, "label": valor} for valor in valores]}


@router.get("/maestro-patrimonial/calidad", response_class=HTMLResponse)
def calidad_maestro_patrimonial(
    request: Request,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    resumen = _resumen_calidad_datos(db)
    return templates.TemplateResponse("maestro_patrimonial_calidad.html", {
        "request": request,
        **resumen,
    })


@router.post("/maestro-patrimonial/cargas/validar")
def validar_archivo_patrimonial(
    request: Request,
    background_tasks: BackgroundTasks,
    archivo: UploadFile = File(...),
    tipo_carga: str = Form("Completa"),
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    nombre = archivo.filename or "reporte.xlsx"
    if not nombre.lower().endswith(".xlsx"):
        return RedirectResponse(
            "/maestro-patrimonial?error=El+reporte+debe+ser+un+archivo+XLSX.",
            status_code=303,
        )
    ruta = ""
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as temporal:
            shutil.copyfileobj(archivo.file, temporal)
            ruta = temporal.name
        carga = preparar_carga(
            db, ruta, nombre, _usuario(request), tipo_carga=tipo_carga
        )
        os.remove(ruta)
        ruta = ""
        background_tasks.add_task(validar_carga_patrimonial_desde_bd, carga.id)
        return RedirectResponse(
            f"/maestro-patrimonial/cargas/{carga.id}", status_code=303
        )
    except Exception as exc:
        if ruta and os.path.exists(ruta):
            os.remove(ruta)
        return RedirectResponse(
            "/maestro-patrimonial?error=" + urlencode({"e": str(exc)})[2:],
            status_code=303,
        )


@router.get("/maestro-patrimonial/cargas/{carga_id}", response_class=HTMLResponse)
def detalle_carga_patrimonial(
    carga_id: int,
    request: Request,
    clasificacion: str = "",
    pagina: int = 1,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    carga = db.get(CargaPatrimonial, carga_id)
    if carga is None:
        raise HTTPException(status_code=404, detail="Carga no encontrada.")
    consulta = db.query(BienCargaPatrimonial).filter(
        BienCargaPatrimonial.carga_id == carga.id
    )
    if clasificacion:
        consulta = consulta.filter(BienCargaPatrimonial.clasificacion == clasificacion)
    total = consulta.count()
    total_paginas = max(1, math.ceil(total / FILAS_CARGA_POR_PAGINA))
    pagina = max(1, min(pagina, total_paginas))
    filas = consulta.order_by(
        BienCargaPatrimonial.codigo_patrimonial
    ).offset((pagina - 1) * FILAS_CARGA_POR_PAGINA).limit(
        FILAS_CARGA_POR_PAGINA
    ).all()
    inicio, fin = rango_registros(pagina, FILAS_CARGA_POR_PAGINA, total)
    base = f"/maestro-patrimonial/cargas/{carga.id}?"
    if clasificacion:
        base += urlencode({"clasificacion": clasificacion})
    errores = json.loads(carga.detalle_validacion or "[]")
    alertas = json.loads(carga.detalle_alertas or "[]")
    return templates.TemplateResponse("maestro_patrimonial_carga.html", {
        "request": request,
        "carga": carga,
        "filas": filas,
        "errores": errores,
        "alertas": alertas,
        "clasificacion": clasificacion,
        "pagina": pagina,
        "total_paginas": total_paginas,
        "paginas": paginas_visibles(pagina, total_paginas),
        "inicio": inicio,
        "fin": fin,
        "total": total,
        "url_base": base,
    })


@router.get("/maestro-patrimonial/cargas/{carga_id}/estado")
def estado_carga_patrimonial(
    carga_id: int,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    carga = db.get(CargaPatrimonial, carga_id)
    if carga is None:
        raise HTTPException(status_code=404, detail="Carga no encontrada.")
    return JSONResponse({
        "estado": carga.estado,
        "progreso": carga.progreso,
        "mensaje": carga.mensaje_progreso or "",
    })


@router.post("/maestro-patrimonial/cargas/{carga_id}/confirmar")
def confirmar_archivo_patrimonial(
    carga_id: int,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    carga = db.get(CargaPatrimonial, carga_id)
    if carga is None or carga.estado not in ("Lista para confirmar", "Interrumpida"):
        raise HTTPException(status_code=400, detail="La carga no está lista para confirmar.")
    if carga.estado == "Interrumpida":
        carga.detalle_validacion = "[]"
    carga.estado = "Procesando"
    db.commit()
    background_tasks.add_task(confirmar_carga_patrimonial, carga.id)
    return RedirectResponse(f"/maestro-patrimonial/cargas/{carga.id}", status_code=303)


@router.post("/maestro-patrimonial/cargas/{carga_id}/cancelar")
def cancelar_archivo_patrimonial(
    carga_id: int,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    carga = db.get(CargaPatrimonial, carga_id)
    if carga is None or carga.estado != "Lista para confirmar":
        raise HTTPException(status_code=400, detail="La carga no se puede cancelar.")
    db.query(BienCargaPatrimonial).filter(
        BienCargaPatrimonial.carga_id == carga.id
    ).delete(synchronize_session=False)
    carga.estado = "Cancelada"
    db.commit()
    return RedirectResponse(f"/maestro-patrimonial/cargas/{carga.id}", status_code=303)


@router.get("/maestro-patrimonial/bienes/{bien_id}", response_class=HTMLResponse)
def detalle_bien_patrimonial(
    bien_id: int,
    request: Request,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    bien = db.query(BienPatrimonial).options(
        selectinload(BienPatrimonial.versiones).selectinload(VersionBienPatrimonial.cambios),
        selectinload(BienPatrimonial.cargas),
        selectinload(BienPatrimonial.correcciones),
        selectinload(BienPatrimonial.conflictos),
    ).filter(BienPatrimonial.id == bien_id).first()
    if bien is None:
        raise HTTPException(status_code=404, detail="Bien no encontrado.")
    (
        altas_detalle,
        impresiones,
        inconsistencias_qr,
        qr_repetido_maestro,
    ) = _relaciones_bien_patrimonial(db, bien)
    versiones = sorted(bien.versiones, key=lambda item: item.creado_en, reverse=True)
    participaciones = sorted(
        bien.cargas,
        key=lambda item: item.carga.creado_en if item.carga else datetime.min,
        reverse=True,
    )
    versiones_por_carga = {
        version.carga_id: version
        for version in versiones if version.carga_id is not None
    }
    linea_tiempo = []

    def agregar_evento(fecha, titulo: str, detalle: str, tipo: str, url: str = ""):
        solo_fecha = isinstance(fecha, date) and not isinstance(fecha, datetime)
        fecha_hora = _fecha_hora_evento(fecha)
        if fecha_hora is not None:
            linea_tiempo.append({
                "fecha": fecha_hora,
                "solo_fecha": solo_fecha,
                "titulo": titulo,
                "detalle": detalle,
                "tipo": tipo,
                "url": url,
            })

    for item in participaciones:
        if item.carga is None:
            continue
        version = versiones_por_carga.get(item.carga_id)
        detalle = f"Resultado: {item.clasificacion}. Carga {item.carga.tipo_carga.lower()}."
        if version and version.cambios:
            detalle += f" Se modificaron {len(version.cambios)} campo(s)."
        agregar_evento(
            item.carga.confirmado_en or item.carga.creado_en,
            f"Carga patrimonial #{item.carga_id}",
            detalle,
            "Carga SIGA",
            f"/maestro-patrimonial/cargas/{item.carga_id}",
        )

    for version in versiones:
        if version.carga_id is not None:
            continue
        detalle = (
            f"{len(version.cambios)} campo(s) modificados por {version.usuario}."
            if version.cambios else f"Registrado por {version.usuario}."
        )
        agregar_evento(
            version.creado_en, version.origen, detalle, "Corrección manual"
        )

    for alta in altas_detalle:
        registro = alta["bien"]
        pecosa = alta["pecosa"]
        referencia = f"Pecosa {pecosa.numero}" if pecosa else "Sin pecosa"
        if alta["expediente_alta"]:
            referencia += f". Expediente de alta {alta['expediente_alta']}"
        agregar_evento(
            registro.fecha_alta,
            "Alta registrada en SIGA",
            referencia,
            "Altas",
        )
        if pecosa:
            agregar_evento(
                pecosa.fecha_recepcion,
                "Pecosa recibida",
                referencia,
                "Altas",
            )
            firma = f"Pecosa {pecosa.numero}"
            if alta["firmante_nombre"]:
                firma += f". Firmante: {alta['firmante_nombre']}"
            if alta["expediente_firma"]:
                firma += f". Expediente de firma {alta['expediente_firma']}"
            agregar_evento(
                pecosa.fecha_firma,
                "Pecosa firmada",
                firma,
                "Altas",
            )

    for item in impresiones:
        inventario = item.inventario
        referencia = (
            inventario.nombre if inventario else f"Inventario #{item.inventario_id}"
        )
        agregar_evento(
            item.importado_en,
            "Incluido en Control Impresión",
            referencia,
            "Control Impresión",
        )
        agregar_evento(
            item.sticker_generado_en,
            "Sticker generado",
            referencia,
            "Control Impresión",
        )
        agregar_evento(
            item.impreso_en,
            "Impresión confirmada",
            referencia,
            "Control Impresión",
        )
    linea_tiempo.sort(key=lambda evento: evento["fecha"], reverse=True)

    return templates.TemplateResponse("maestro_patrimonial_detalle.html", {
        "request": request,
        "bien": bien,
        "campos": [(campo, ETIQUETAS_CAMPOS[campo]) for campo in CAMPOS],
        "campos_editables": [(campo, ETIQUETAS_CAMPOS[campo], CAMPOS[campo][1]) for campo in CAMPOS_EDITABLES],
        "campos_opcionales": CAMPOS_OPCIONALES,
        "versiones": versiones,
        "participaciones": participaciones,
        "linea_tiempo": linea_tiempo,
        "altas": altas_detalle,
        "impresiones": impresiones,
        "inconsistencias_qr": inconsistencias_qr,
        "qr_repetido_maestro": qr_repetido_maestro,
        "conflictos": [c for c in bien.conflictos if c.estado == "Pendiente"],
        "etiquetas": ETIQUETAS_CAMPOS,
    })


def _relaciones_bien_patrimonial(db: Session, bien: BienPatrimonial):
    altas = db.query(BienAlta).options(
        selectinload(BienAlta.pecosa).selectinload(Pecosa.expediente)
    ).filter(
        BienAlta.codigo_patrimonial == bien.codigo_patrimonial
    ).all()
    altas_detalle = []
    for alta in altas:
        pecosa = alta.pecosa
        nombre_firmante, dni_firmante = _datos_firmante(
            db, pecosa.firmante if pecosa else ""
        )
        altas_detalle.append({
            "bien": alta,
            "pecosa": pecosa,
            "expediente_alta": (
                pecosa.expediente.numero
                if pecosa and pecosa.expediente else ""
            ),
            "expediente_firma": pecosa.expediente_firma if pecosa else "",
            "firmante_nombre": nombre_firmante,
            "firmante_dni": dni_firmante,
        })
    impresiones = db.query(BienInventarioImpresion).filter(
        BienInventarioImpresion.codigo_patrimonial == bien.codigo_patrimonial
    ).options(
        selectinload(BienInventarioImpresion.inventario)
    ).all()
    inconsistencias_qr = []
    qr_repetido_maestro = []
    if bien.codigo_qr:
        qr_repetido_maestro = db.query(BienPatrimonial).filter(
            BienPatrimonial.codigo_qr == bien.codigo_qr,
            BienPatrimonial.id != bien.id,
        ).order_by(BienPatrimonial.codigo_patrimonial).all()
        inconsistencias_qr.extend([
            f"Alta: {fila.codigo_patrimonial}"
            for fila in db.query(BienAlta).filter(
                BienAlta.codigo_qr == bien.codigo_qr,
                BienAlta.codigo_patrimonial != bien.codigo_patrimonial,
            ).all()
        ])
        inconsistencias_qr.extend([
            f"Control Impresión: {fila.codigo_patrimonial}"
            for fila in db.query(BienInventarioImpresion).filter(
                BienInventarioImpresion.codigo_qr == bien.codigo_qr,
                BienInventarioImpresion.codigo_patrimonial != bien.codigo_patrimonial,
            ).all()
        ])
    return altas_detalle, impresiones, inconsistencias_qr, qr_repetido_maestro


@router.get("/maestro-patrimonial/bienes/{bien_id}/ficha.pdf")
def ficha_pdf_bien_patrimonial(
    bien_id: int,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    bien = db.query(BienPatrimonial).options(
        selectinload(BienPatrimonial.versiones).selectinload(
            VersionBienPatrimonial.cambios
        ),
        selectinload(BienPatrimonial.conflictos),
    ).filter(BienPatrimonial.id == bien_id).first()
    if bien is None:
        raise HTTPException(status_code=404, detail="Bien no encontrado.")
    altas, impresiones, inconsistencias_qr, qr_repetido = (
        _relaciones_bien_patrimonial(db, bien)
    )
    contenido = generar_ficha_activo_pdf(
        bien,
        altas,
        impresiones,
        sorted(bien.versiones, key=lambda item: item.creado_en, reverse=True),
        inconsistencias_qr=inconsistencias_qr,
        qr_repetido_maestro=qr_repetido,
        conflictos=bien.conflictos,
    )
    nombre = re.sub(r"[^A-Za-z0-9_-]+", "_", bien.codigo_patrimonial)
    return Response(
        content=contenido,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'inline; filename="ficha_activo_{nombre}.pdf"'
        },
    )


@router.post("/maestro-patrimonial/bienes/{bien_id}/editar")
async def editar_bien(
    bien_id: int,
    request: Request,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    bien = db.get(BienPatrimonial, bien_id)
    if bien is None:
        raise HTTPException(status_code=404, detail="Bien no encontrado.")
    formulario = await request.form()
    valores = {campo: formulario.get(campo, "") for campo in CAMPOS_EDITABLES}
    try:
        cantidad = editar_bien_patrimonial(
            db, bien, valores, formulario.get("motivo", ""), _usuario(request)
        )
    except ValueError as exc:
        return RedirectResponse(
            f"/maestro-patrimonial/bienes/{bien.id}?error="
            + urlencode({"e": str(exc)})[2:],
            status_code=303,
        )
    return RedirectResponse(
        f"/maestro-patrimonial/bienes/{bien.id}?info="
        + urlencode({"i": f"Se registraron {cantidad} cambios."})[2:],
        status_code=303,
    )


@router.post("/maestro-patrimonial/conflictos/{conflicto_id}/resolver")
def resolver_conflicto_bien(
    conflicto_id: int,
    request: Request,
    decision: str = Form(...),
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    conflicto = db.get(ConflictoBienPatrimonial, conflicto_id)
    if conflicto is None:
        raise HTTPException(status_code=404, detail="Conflicto no encontrado.")
    bien_id = conflicto.bien_id
    try:
        resolver_conflicto(db, conflicto, decision, _usuario(request))
    except ValueError as exc:
        return RedirectResponse(
            f"/maestro-patrimonial/bienes/{bien_id}?error="
            + urlencode({"e": str(exc)})[2:],
            status_code=303,
        )
    return RedirectResponse(
        f"/maestro-patrimonial/bienes/{bien_id}?info=Conflicto+resuelto.",
        status_code=303,
    )


def _valor_exportable(valor):
    if isinstance(valor, Decimal):
        return float(valor)
    return valor


def _ancho_columna_exportacion(encabezado: str) -> float:
    """Asigna un ancho legible sin permitir columnas excesivamente grandes."""
    nombre = unicodedata.normalize("NFKD", encabezado.casefold())
    nombre = "".join(
        caracter for caracter in nombre
        if not unicodedata.combining(caracter)
    )
    if "fecha" in nombre:
        return 14
    if "valor" in nombre or "precio" in nombre:
        return 16
    if nombre in {"descripcion", "caracteristicas", "observaciones"}:
        return 45
    if any(texto in nombre for texto in (
        "depend", "responsable", "usuario", "ubicac", "nombre_item",
    )):
        return 36
    if any(texto in nombre for texto in (
        "codigo", "nro_", "numero", "secuencia", "modelo", "serie",
    )):
        return 22
    return min(max(len(encabezado) + 3, 13), 28)


def _configurar_hoja_exportacion(
    hoja, total_filas: int, encabezados: list[str]
) -> list[float]:
    hoja.freeze_panes = "A2"
    ultima_columna = get_column_letter(len(encabezados))
    hoja.auto_filter.ref = f"A1:{ultima_columna}{total_filas + 1}"
    hoja.sheet_view.showGridLines = False
    hoja.row_dimensions[1].height = 32
    anchos = [_ancho_columna_exportacion(encabezado) for encabezado in encabezados]
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    return anchos


def _agregar_encabezado_excel(hoja, encabezados: list[str]) -> None:
    fila = []
    for titulo in encabezados:
        celda = WriteOnlyCell(hoja, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")
        celda.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        fila.append(celda)
    hoja.append(fila)


def _celda_excel(hoja, valor, tipo: str, ancho: float):
    if valor is None:
        return None
    if tipo == "fecha":
        celda = WriteOnlyCell(hoja, value=valor)
        celda.number_format = "dd/mm/yyyy"
        return celda
    if tipo == "decimal":
        celda = WriteOnlyCell(hoja, value=_valor_exportable(valor))
        celda.number_format = "#,##0.00"
        return celda
    if isinstance(valor, str) and len(valor) > max(12, int(ancho) - 2):
        celda = WriteOnlyCell(hoja, value=valor)
        celda.alignment = Alignment(vertical="top", wrap_text=True)
        return celda
    return valor


def _fila_siga_exportable(bien: BienPatrimonial) -> list:
    valores = json.loads(bien.datos_fuente)
    if len(valores) < len(ENCABEZADOS_SIGA):
        valores.extend([None] * (len(ENCABEZADOS_SIGA) - len(valores)))
    for campo, (indice, _) in CAMPOS.items():
        valores[indice] = getattr(bien, campo)
    return valores[:len(ENCABEZADOS_SIGA)]


def _actualizar_exportacion(
    db: Session,
    exportacion: ExportacionPatrimonial,
    progreso: int,
) -> None:
    exportacion.progreso = progreso
    db.commit()


def generar_exportacion_patrimonial(exportacion_id: int) -> None:
    """Genera el Excel fuera de la solicitud web y conserva el resultado."""
    db = SessionLocal()
    try:
        exportacion = db.get(ExportacionPatrimonial, exportacion_id)
        if exportacion is None or exportacion.estado not in ("Pendiente", "Procesando"):
            return
        exportacion.estado = "Procesando"
        exportacion.progreso = 2
        exportacion.mensaje_error = None
        db.commit()

        filtros = json.loads(exportacion.filtros or "{}")
        campos_resumen = json.loads(exportacion.columnas or "[]")
        consulta = _aplicar_filtros(db.query(BienPatrimonial), filtros)
        total = consulta.order_by(None).count()
        exportacion.total_filas = total
        _actualizar_exportacion(db, exportacion, 5)

        if exportacion.tipo_reporte == "SIGA completo":
            encabezados = list(ENCABEZADOS_SIGA)
            tipos = [TIPOS_COLUMNAS_EXPORTACION.get(i, "") for i in range(len(encabezados))]
        else:
            campos_resumen = [campo for campo in campos_resumen if campo in CAMPOS]
            if not campos_resumen:
                campos_resumen = [
                    "codigo_patrimonial", "codigo_qr", "descripcion",
                    "nombre_dependencia", "usuario", "ubicacion_fisica",
                    "marca", "modelo", "estado_conservacion",
                ]
            encabezados = [ETIQUETAS_CAMPOS[campo] for campo in campos_resumen]
            tipos = [CAMPOS[campo][1] for campo in campos_resumen]

        libro = Workbook(write_only=True)
        resumen = libro.create_sheet("Resumen")
        resumen.column_dimensions["A"].width = 24
        resumen.column_dimensions["B"].width = 45
        resumen.append(["Reporte", exportacion.tipo_reporte])
        resumen.append(["Generado", datetime.now()])
        resumen.append(["Usuario", exportacion.usuario])
        resumen.append(["Cantidad de bienes", total])
        resumen.append([])
        resumen.append(["Filtro", "Valor"])
        for clave, valor in filtros.items():
            if valor:
                resumen.append([clave, valor])

        hoja = libro.create_sheet("Bienes")
        anchos = _configurar_hoja_exportacion(hoja, total, encabezados)
        _agregar_encabezado_excel(hoja, encabezados)

        procesados = 0
        ultimo_codigo = ""
        while True:
            pagina = consulta.filter(
                BienPatrimonial.codigo_patrimonial > ultimo_codigo
            ).order_by(BienPatrimonial.codigo_patrimonial).limit(500).all()
            if not pagina:
                break
            for bien in pagina:
                if exportacion.tipo_reporte == "SIGA completo":
                    valores = _fila_siga_exportable(bien)
                else:
                    valores = [getattr(bien, campo) for campo in campos_resumen]
                hoja.append([
                    _celda_excel(hoja, valor, tipos[indice], anchos[indice])
                    for indice, valor in enumerate(valores)
                ])
            procesados += len(pagina)
            ultimo_codigo = pagina[-1].codigo_patrimonial
            progreso = 90 if total == 0 else min(90, 5 + int(procesados * 85 / total))
            _actualizar_exportacion(db, exportacion, progreso)

        salida = io.BytesIO()
        libro.save(salida)
        marca_tiempo = datetime.now().strftime("%Y%m%d_%H%M%S")
        exportacion.nombre_archivo = f"maestro_patrimonial_{marca_tiempo}.xlsx"
        exportacion.archivo_contenido = salida.getvalue()
        exportacion.estado = "Completada"
        exportacion.progreso = 100
        exportacion.completado_en = datetime.utcnow()
        db.commit()
    except Exception as exc:
        db.rollback()
        exportacion = db.get(ExportacionPatrimonial, exportacion_id)
        if exportacion:
            exportacion.estado = "Error"
            exportacion.progreso = 0
            exportacion.mensaje_error = (
                "No se pudo generar el archivo. "
                f"Tipo de error: {type(exc).__name__}."
            )
            db.commit()
    finally:
        db.close()


@router.get("/maestro-patrimonial/exportar")
def configurar_exportacion_patrimonial(
    request: Request,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    filtros = _parametros_filtros(request)
    total = _aplicar_filtros(db.query(BienPatrimonial), filtros).count()
    campos = [
        (campo, ETIQUETAS_CAMPOS[campo])
        for campo in CAMPOS
    ]
    predeterminados = {
        "codigo_patrimonial", "codigo_qr", "descripcion", "nombre_dependencia",
        "usuario", "ubicacion_fisica", "marca", "modelo", "estado_conservacion",
    }
    return templates.TemplateResponse("maestro_patrimonial_exportar.html", {
        "request": request,
        "filtros": filtros,
        "filtros_json": json.dumps(filtros, ensure_ascii=False),
        "total": total,
        "campos": campos,
        "predeterminados": predeterminados,
    })


@router.post("/maestro-patrimonial/exportaciones")
async def crear_exportacion_patrimonial(
    request: Request,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    formulario = await request.form()
    tipo_reporte = str(formulario.get("tipo_reporte", "Resumen personalizado"))
    if tipo_reporte not in ("Resumen personalizado", "SIGA completo"):
        raise HTTPException(status_code=400, detail="Tipo de reporte no válido.")
    try:
        filtros_recibidos = json.loads(str(formulario.get("filtros", "{}")))
    except (TypeError, ValueError):
        filtros_recibidos = {}
    filtros = {
        nombre: str(filtros_recibidos.get(nombre, "")).strip()
        for nombre in [
            "codigo_patrimonial", "codigo_qr", "descripcion", "q",
            *FILTROS_TEXTO,
            *(nombre for campo in FILTROS_FECHA for nombre in (
                f"{campo}_desde", f"{campo}_hasta"
            )),
        ]
    }
    columnas = [
        str(valor) for valor in formulario.getlist("columnas")
        if str(valor) in CAMPOS
    ]
    exportacion = ExportacionPatrimonial(
        usuario=_usuario(request),
        estado="Pendiente",
        progreso=0,
        tipo_reporte=tipo_reporte,
        columnas=json.dumps(columnas, ensure_ascii=False),
        filtros=json.dumps(filtros, ensure_ascii=False),
    )
    db.add(exportacion)
    db.commit()
    db.refresh(exportacion)
    background_tasks.add_task(generar_exportacion_patrimonial, exportacion.id)
    return RedirectResponse(
        f"/maestro-patrimonial/exportaciones/{exportacion.id}", status_code=303
    )


@router.get("/maestro-patrimonial/exportaciones/{exportacion_id}", response_class=HTMLResponse)
def detalle_exportacion_patrimonial(
    exportacion_id: int,
    request: Request,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    exportacion = db.get(ExportacionPatrimonial, exportacion_id)
    if exportacion is None:
        raise HTTPException(status_code=404, detail="Exportación no encontrada.")
    return templates.TemplateResponse("maestro_patrimonial_exportacion.html", {
        "request": request,
        "exportacion": exportacion,
    })


@router.get("/maestro-patrimonial/exportaciones/{exportacion_id}/estado")
def estado_exportacion_patrimonial(
    exportacion_id: int,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    exportacion = db.get(ExportacionPatrimonial, exportacion_id)
    if exportacion is None:
        raise HTTPException(status_code=404, detail="Exportación no encontrada.")
    return JSONResponse({
        "estado": exportacion.estado,
        "progreso": exportacion.progreso,
        "error": exportacion.mensaje_error or "",
    })


@router.get("/maestro-patrimonial/exportaciones/{exportacion_id}/descargar")
def descargar_exportacion_patrimonial(
    exportacion_id: int,
    db: Session = Depends(get_db),
    _=Depends(requiere_login),
):
    exportacion = db.get(ExportacionPatrimonial, exportacion_id)
    if (
        exportacion is None or exportacion.estado != "Completada"
        or not exportacion.archivo_contenido
    ):
        raise HTTPException(status_code=404, detail="El archivo aún no está disponible.")
    return Response(
        content=exportacion.archivo_contenido,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{exportacion.nombre_archivo}"'
            )
        },
    )


def reanudar_tareas_patrimoniales() -> None:
    """Reanuda cargas y exportaciones que quedaron activas tras un reinicio."""
    db = SessionLocal()
    try:
        validaciones = [
            carga_id for (carga_id,) in db.query(CargaPatrimonial.id).filter(
                CargaPatrimonial.estado == "Validando"
            ).all()
        ]
        confirmaciones = [
            carga_id for (carga_id,) in db.query(CargaPatrimonial.id).filter(
                CargaPatrimonial.estado == "Procesando"
            ).all()
        ]
        exportaciones = [
            exportacion_id
            for (exportacion_id,) in db.query(ExportacionPatrimonial.id).filter(
                ExportacionPatrimonial.estado.in_(("Pendiente", "Procesando"))
            ).all()
        ]
    finally:
        db.close()
    for carga_id in validaciones:
        threading.Thread(
            target=validar_carga_patrimonial_desde_bd,
            args=(carga_id,), daemon=True,
        ).start()
    for carga_id in confirmaciones:
        threading.Thread(
            target=confirmar_carga_patrimonial,
            args=(carga_id,), daemon=True,
        ).start()
    for exportacion_id in exportaciones:
        threading.Thread(
            target=generar_exportacion_patrimonial,
            args=(exportacion_id,), daemon=True,
        ).start()
