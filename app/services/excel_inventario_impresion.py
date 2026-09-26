"""Lectura y actualización del reporte general de impresión de One Vision."""
from datetime import datetime
from pathlib import Path
import re
import unicodedata
from urllib.parse import urlsplit

from openpyxl import load_workbook
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.models import (
    BienAlta,
    BienInventarioImpresion,
    CargaInventarioImpresion,
    InventarioImpresion,
)
from app.services.solicitudes_impresion import reconciliar_solicitudes_pendientes


COLUMNAS = {
    "codigo_patrimonial": ("codigo patrimonial",),
    "codigo_qr": ("codigo qr",),
    "ruta_qr": ("ruta qr", "ruta", "url"),
    "descripcion": ("bien", "descripcion"),
    "establecimiento": ("establecimiento",),
    "red": ("red",),
    "area": ("area",),
    "marca": ("marca",),
    "modelo": ("modelo",),
    "color": ("color",),
    "nro_serie": ("nr serie", "nro serie", "numero serie"),
}
REQUERIDAS = (
    "codigo_patrimonial", "codigo_qr", "ruta_qr", "descripcion",
    "establecimiento", "red", "area",
)
TAMANO_LOTE_BD = 1000


def _normalizar_encabezado(valor) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"[^a-z0-9]+", " ", texto.casefold())
    return " ".join(texto.split())


def _texto_excel(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return " ".join(str(valor).strip().split())


def normalizar_codigo_patrimonial(valor) -> str:
    texto = _texto_excel(valor)
    if texto.isdigit() and len(texto) == 11:
        return texto.zfill(12)
    return texto


def clasificar_tipo_bien(codigo_patrimonial: str) -> str:
    """Clasifica el bien después de recuperar el cero inicial de SIGA."""
    return (
        "Activo fijo"
        if codigo_patrimonial and len(codigo_patrimonial) >= 12
        else "Sobrante"
    )


def normalizar_codigo_qr(valor) -> str:
    texto = _texto_excel(valor)
    if texto.isdigit():
        return texto.lstrip("0") or "0"
    return texto


def _indices_columnas(encabezados) -> dict[str, int]:
    normalizados = [_normalizar_encabezado(valor) for valor in encabezados]
    indices = {}
    for clave, candidatos in COLUMNAS.items():
        for indice, encabezado in enumerate(normalizados):
            if encabezado in candidatos:
                indices[clave] = indice
                break
    faltantes = [clave for clave in REQUERIDAS if clave not in indices]
    if faltantes:
        raise ValueError(
            "El reporte no contiene las columnas requeridas: "
            + ", ".join(faltantes)
        )
    return indices


def _motivo_bloqueo(codigo_qr: str, ruta_qr: str) -> str | None:
    if not codigo_qr:
        return "Sin Código QR"
    if not ruta_qr:
        return "Sin Ruta QR"
    partes = urlsplit(ruta_qr)
    if partes.scheme.lower() not in ("http", "https") or not partes.netloc:
        return "Ruta QR no válida"
    return None


def iterar_reporte_inventario(ruta_archivo: str):
    """Itera el Excel en modo lectura para no cargar el libro completo."""
    libro = load_workbook(ruta_archivo, read_only=True, data_only=True)
    try:
        hoja = libro.active
        filas = hoja.iter_rows(values_only=True)
        encabezados = next(filas, None)
        if encabezados is None:
            raise ValueError("El reporte está vacío.")
        indices = _indices_columnas(encabezados)

        for numero_fila, fila in enumerate(filas, start=2):
            def valor(clave):
                indice = indices.get(clave)
                return _texto_excel(fila[indice]) if indice is not None else ""

            codigo = normalizar_codigo_patrimonial(valor("codigo_patrimonial"))
            codigo_qr = normalizar_codigo_qr(valor("codigo_qr"))
            if not codigo and not codigo_qr:
                yield numero_fila, None
                continue
            tipo_bien = clasificar_tipo_bien(codigo)
            ruta_qr = valor("ruta_qr")
            motivo = _motivo_bloqueo(codigo_qr, ruta_qr)
            yield numero_fila, {
                # Un código corto describe la clase del sobrante, pero no es un
                # identificador patrimonial único. El QR identifica cada bien.
                "codigo_patrimonial": (
                    codigo if tipo_bien == "Activo fijo" else None
                ),
                "codigo_qr": codigo_qr or None,
                "tipo_bien": tipo_bien,
                "ruta_qr": ruta_qr or None,
                "descripcion": valor("descripcion"),
                "establecimiento": valor("establecimiento") or None,
                "red": valor("red") or None,
                "area": valor("area") or None,
                "marca": valor("marca") or None,
                "modelo": valor("modelo") or None,
                "color": valor("color") or None,
                "nro_serie": valor("nro_serie") or None,
                "imprimible": 0 if motivo else 1,
                "motivo_bloqueo": motivo,
            }
    finally:
        libro.close()


def importar_reporte_inventario(
    db: Session,
    ruta_archivo: str,
    nombre_archivo: str,
    anio: str = "2026",
    unidad_ejecutora: str = "DIRESA",
) -> dict:
    """Actualiza el inventario anual sin borrar lotes ni impresiones previas."""
    registros = []
    codigos_vistos = set()
    duplicados = 0
    omitidos = 0
    for _, registro in iterar_reporte_inventario(ruta_archivo):
        if registro is None:
            omitidos += 1
            continue
        clave = (
            ("P", registro["codigo_patrimonial"])
            if registro["tipo_bien"] == "Activo fijo"
            else ("Q", registro["codigo_qr"] or registro["codigo_patrimonial"])
        )
        if clave in codigos_vistos:
            duplicados += 1
            continue
        codigos_vistos.add(clave)
        registros.append(registro)

    if not registros:
        raise ValueError("El reporte no contiene bienes con código patrimonial ni QR.")

    inventario = (
        db.query(InventarioImpresion)
        .filter(
            InventarioImpresion.anio == str(anio),
            InventarioImpresion.unidad_ejecutora == unidad_ejecutora,
        )
        .first()
    )
    if inventario is None:
        inventario = InventarioImpresion(
            nombre=f"Inventario UE {unidad_ejecutora} {anio}",
            anio=str(anio),
            unidad_ejecutora=unidad_ejecutora,
        )
        db.add(inventario)
        db.flush()

    reactivados = db.query(BienInventarioImpresion).filter(
        BienInventarioImpresion.inventario_id == inventario.id,
        BienInventarioImpresion.activo != 1,
    ).update({BienInventarioImpresion.activo: 1}, synchronize_session=False)

    bienes_existentes = db.query(BienInventarioImpresion).filter(
        BienInventarioImpresion.inventario_id == inventario.id
    ).all()
    existentes_por_patrimonial = {}
    for bien in bienes_existentes:
        codigo_existente = normalizar_codigo_patrimonial(bien.codigo_patrimonial)
        if codigo_existente:
            existentes_por_patrimonial.setdefault(codigo_existente, []).append(bien)
    existentes_por_qr = {}
    for bien in bienes_existentes:
        if bien.codigo_qr:
            existentes_por_qr.setdefault(bien.codigo_qr, []).append(bien)

    altas_por_codigo: dict[str, list[int]] = {}
    for alta_id, codigo in db.query(BienAlta.id, BienAlta.codigo_patrimonial).all():
        normalizado = normalizar_codigo_patrimonial(codigo)
        if normalizado:
            altas_por_codigo.setdefault(normalizado, []).append(alta_id)

    ahora = datetime.utcnow()
    nuevos = []
    actualizaciones = []
    sin_cambios = 0
    vinculados = 0
    ambiguos = 0
    activos_fijos = 0
    sobrantes = 0
    for registro in registros:
        codigo_patrimonial = registro["codigo_patrimonial"]
        es_activo_fijo = registro["tipo_bien"] == "Activo fijo"
        coincidencias = (
            altas_por_codigo.get(codigo_patrimonial, [])
            if es_activo_fijo and codigo_patrimonial
            else []
        )
        if len(coincidencias) == 1:
            registro["bien_alta_id"] = coincidencias[0]
            registro["relacion_alta"] = "Vinculado"
            vinculados += 1
        elif len(coincidencias) > 1:
            registro["bien_alta_id"] = None
            registro["relacion_alta"] = "Ambiguo"
            ambiguos += 1
        else:
            registro["bien_alta_id"] = None
            registro["relacion_alta"] = "Sin alta"

        if es_activo_fijo:
            activos_fijos += 1
        else:
            sobrantes += 1

        registro.update({
            "inventario_id": inventario.id,
            "activo": 1,
        })
        candidatos_codigo = (
            existentes_por_patrimonial.get(codigo_patrimonial, [])
            if es_activo_fijo and codigo_patrimonial
            else []
        )
        existente = candidatos_codigo[0] if len(candidatos_codigo) == 1 else None
        if existente is None and registro["codigo_qr"]:
            candidatos_qr = existentes_por_qr.get(registro["codigo_qr"], [])
            candidatos_compatibles = [
                bien for bien in candidatos_qr
                if (
                    not es_activo_fijo
                    or not bien.codigo_patrimonial
                    or normalizar_codigo_patrimonial(bien.codigo_patrimonial)
                    == codigo_patrimonial
                )
            ]
            if len(candidatos_compatibles) == 1:
                existente = candidatos_compatibles[0]
        if existente:
            estado_anterior = existente.estado_impresion
            generado_en = existente.sticker_generado_en
            nuevo_estado = estado_anterior
            if estado_anterior != "Impreso":
                if not registro["imprimible"]:
                    nuevo_estado = "Bloqueado"
                elif estado_anterior == "Bloqueado":
                    nuevo_estado = (
                        "Sticker generado" if generado_en else "Pendiente"
                    )
            registro["estado_impresion"] = nuevo_estado
            campos_comparables = (
                "bien_alta_id", "relacion_alta", "codigo_patrimonial",
                "codigo_qr", "tipo_bien", "ruta_qr", "descripcion",
                "establecimiento", "red", "area", "marca", "modelo",
                "color", "nro_serie", "imprimible", "motivo_bloqueo",
                "estado_impresion", "activo",
            )
            cambio = any(
                getattr(existente, campo) != registro.get(campo)
                for campo in campos_comparables
            )
            if cambio:
                actualizaciones.append({
                    "id": existente.id,
                    "actualizado_en": ahora,
                    **registro,
                })
            else:
                sin_cambios += 1
        else:
            nuevos.append({
                "importado_en": ahora,
                "actualizado_en": ahora,
                "estado_impresion": (
                    "Pendiente" if registro["imprimible"] else "Bloqueado"
                ),
                **registro,
            })

    for inicio in range(0, len(actualizaciones), TAMANO_LOTE_BD):
        db.bulk_update_mappings(
            BienInventarioImpresion,
            actualizaciones[inicio:inicio + TAMANO_LOTE_BD],
        )
    for inicio in range(0, len(nuevos), TAMANO_LOTE_BD):
        db.bulk_insert_mappings(
            BienInventarioImpresion,
            nuevos[inicio:inicio + TAMANO_LOTE_BD],
        )

    db.flush()
    solicitudes = reconciliar_solicitudes_pendientes(db, inventario.id)

    totales_tipo = dict(
        db.query(
            BienInventarioImpresion.tipo_bien,
            func.count(BienInventarioImpresion.id),
        ).filter(
            BienInventarioImpresion.inventario_id == inventario.id,
            BienInventarioImpresion.activo == 1,
        ).group_by(BienInventarioImpresion.tipo_bien).all()
    )
    total_universo = sum(totales_tipo.values())

    inventario.archivo_ultimo = Path(nombre_archivo).name
    inventario.total_registros = total_universo
    inventario.actualizado_en = ahora
    carga = CargaInventarioImpresion(
        inventario_id=inventario.id,
        archivo=Path(nombre_archivo).name,
        creado_en=ahora,
        filas_procesadas=len(registros),
        nuevos=len(nuevos),
        actualizados=len(actualizaciones),
        sin_cambios=sin_cambios,
        activos_fijos=activos_fijos,
        sobrantes=sobrantes,
        duplicados=duplicados,
        omitidos=omitidos,
        total_universo=total_universo,
        activos_fijos_universo=totales_tipo.get("Activo fijo", 0),
        sobrantes_universo=totales_tipo.get("Sobrante", 0),
    )
    db.add(carga)
    db.flush()
    db.commit()
    return {
        "carga_id": carga.id,
        "inventario_id": inventario.id,
        "total": len(registros),
        "nuevos": len(nuevos),
        "actualizados": len(actualizaciones),
        "sin_cambios": sin_cambios,
        "duplicados": duplicados,
        "omitidos": omitidos,
        "vinculados": vinculados,
        "ambiguos": ambiguos,
        "inactivos": 0,
        "reactivados": reactivados,
        "activos_fijos": activos_fijos,
        "sobrantes": sobrantes,
        "total_universo": total_universo,
        "activos_fijos_universo": totales_tipo.get("Activo fijo", 0),
        "sobrantes_universo": totales_tipo.get("Sobrante", 0),
        "solicitudes_vinculadas": solicitudes["vinculados"],
        "solicitudes_observadas": solicitudes["observados"],
        "solicitudes_esperando": solicitudes["esperando"],
    }


def reparar_universos_acumulativos(db: Session) -> dict:
    """Recupera bienes ocultos y aplica la clasificación vigente una sola vez."""
    reactivados = db.query(BienInventarioImpresion).filter(
        BienInventarioImpresion.activo != 1
    ).update({BienInventarioImpresion.activo: 1}, synchronize_session=False)

    reclasificados = 0
    normalizados = 0
    candidatos = db.query(BienInventarioImpresion).filter(or_(
        BienInventarioImpresion.codigo_patrimonial.is_(None),
        func.length(BienInventarioImpresion.codigo_patrimonial) < 12,
    )).all()
    for bien in candidatos:
        codigo_normalizado = normalizar_codigo_patrimonial(
            bien.codigo_patrimonial
        )
        tipo_bien = clasificar_tipo_bien(codigo_normalizado)
        if tipo_bien == "Sobrante":
            if bien.codigo_patrimonial is not None:
                bien.codigo_patrimonial = None
                normalizados += 1
        elif codigo_normalizado and codigo_normalizado != bien.codigo_patrimonial:
            conflicto = db.query(BienInventarioImpresion.id).filter(
                BienInventarioImpresion.inventario_id == bien.inventario_id,
                BienInventarioImpresion.codigo_patrimonial == codigo_normalizado,
                BienInventarioImpresion.id != bien.id,
            ).first()
            if conflicto is None:
                bien.codigo_patrimonial = codigo_normalizado
                normalizados += 1
        if bien.tipo_bien != tipo_bien:
            bien.tipo_bien = tipo_bien
            reclasificados += 1

    for inventario in db.query(InventarioImpresion).all():
        inventario.total_registros = db.query(
            func.count(BienInventarioImpresion.id)
        ).filter(
            BienInventarioImpresion.inventario_id == inventario.id,
            BienInventarioImpresion.activo == 1,
        ).scalar() or 0

    return {
        "reactivados": reactivados,
        "normalizados": normalizados,
        "reclasificados": reclasificados,
    }
