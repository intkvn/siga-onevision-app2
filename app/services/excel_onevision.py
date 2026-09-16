"""
Dos cosas:
1) Generar el archivo .xls que pide el Formato de Importación
   de One Visión (16 columnas), a partir de los BienAlta ya normalizados.
2) Leer el reporte que se descarga DESPUÉS de cargar a One Visión
   (con Código QR y Ruta QR), corrigiendo el bug del código patrimonial
   (llega con el primer carácter vacío) para poder cruzarlo.
"""
from collections.abc import Iterator

import xlwt
from openpyxl import load_workbook
from app.config import ESTADOS

ENCABEZADOS = [
    "QR", "AÑO", "Ejecutora", "IPRESS", "DNI", "Codigo Patrimonial", "Descripcion",
    "Fecha de Alta", "Modelo", "Marca", "Estado", "Nro. Serie",
    "Observaciones", "Color", "Observacion Analista", "Caracteristicas",
]

DESCRIPCIONES = [
    "QR del\nEquipo",
    "Año del\nInventario",
    "Codigo de\nla Ejecutora",
    "Codigo IPRESS\ndel Establecimiento",
    "Número de DNI del\nResponsable del Equipo",
    "Código Patrimonial\ndel Equipo",
    "Descripción\ndel Equipo",
    "Fecha de Alta\ndel Equipo",
    "Modelo\ndel Equipo",
    "Marca\ndel Equipo",
    "(Nuevo, Bueno, Regular,\nMalo, Muy Malo,\nChatarra, RAEE)",
    "Nro. de Serie\ndel Equipo",
    "Observaciones\ndel Equipo",
    "Color\ndel Equipo",
    "Observaciones\ndel analista",
    "Caracteristicas\ndel Equipo",
]

REQUERIDOS = [
    "Opcional", "Obligatorio(*)", "Obligatorio(*)", "Obligatorio(*)",
    "Opcional", "Opcional", "Opcional", "Opcional", "Opcional", "Opcional",
    "Opcional", "Opcional", "Opcional", "Opcional", "Opcional", "Opcional",
]


def _estilos_formato_onevision():
    """Estilos visuales definidos por el formato de importación de One Visión."""
    borde = "borders: left thin, right thin, top thin, bottom thin;"
    azul = "pattern: pattern solid, fore_colour 0x21;"
    amarillo = "pattern: pattern solid, fore_colour 0x22;"
    comun = "font: name Calibri;"

    return {
        "titulo": xlwt.easyxf(
            f"{comun} font: bold on, height 400, colour white; {azul} "
            "align: horiz center, vert bottom, wrap on;"
        ),
        "subtitulo": xlwt.easyxf(
            f"{comun} font: bold on, height 320; {amarillo} "
            "align: horiz center, vert bottom;"
        ),
        "nota": xlwt.easyxf(
            f"{comun} font: bold on, height 220; align: vert bottom, wrap on;"
        ),
        "cabecera": xlwt.easyxf(
            f"{comun} font: bold on, height 240, colour white; {azul} {borde} "
            "align: horiz center, vert bottom;"
        ),
        "descripcion": xlwt.easyxf(
            f"{comun} font: bold on, height 200; {amarillo} {borde} "
            "align: horiz center, vert bottom, wrap on;"
        ),
        "opcional": xlwt.easyxf(
            f"{comun} font: bold on, height 200, colour 0x23; {amarillo} {borde} "
            "align: horiz center, vert bottom;"
        ),
        "obligatorio": xlwt.easyxf(
            f"{comun} font: bold on, height 200, colour 0x24; {amarillo} {borde} "
            "align: horiz center, vert bottom;"
        ),
        "dato": xlwt.easyxf(
            f"{comun} {borde} align: vert bottom;"
        ),
    }


def generar_formato_importacion(bienes: list, anio: str, ejecutora: str, ruta_salida: str):
    """
    'bienes' es una lista de objetos BienAlta ya con persona y centro_costo
    asignados (el cruce ya resuelto, sin pendientes).
    Escribe el archivo .xls en 'ruta_salida'.
    """
    libro = xlwt.Workbook(encoding="utf-8")
    libro.set_colour_RGB(0x21, 30, 136, 229)  # azul del formato de One Visión
    libro.set_colour_RGB(0x22, 251, 247, 205)  # amarillo del formato de One Visión
    libro.set_colour_RGB(0x23, 4, 134, 190)    # texto de campos opcionales
    libro.set_colour_RGB(0x24, 234, 67, 53)    # texto de campos obligatorios
    hoja = libro.add_sheet("Worksheet")
    estilos = _estilos_formato_onevision()

    # Presentación institucional del formato entregado por One Visión.
    hoja.write(0, 0, 1)
    hoja.write(0, 1, 2)
    hoja.write(0, 2, 3)
    hoja.merge(1, 1, 0, 14)
    hoja.write(1, 0, "ONE VISION - EQUIPAMIENTO\n", estilos["titulo"])
    hoja.merge(2, 2, 0, 15)
    hoja.write(2, 0, "FORMATO DE IMPORTACIÓN", estilos["subtitulo"])
    hoja.merge(3, 3, 0, 15)
    hoja.write(
        3,
        0,
        "Observaciones: Solo llenar los campos solicitados en la parte inferior,\n"
        "no modificar el formato ya que tiene parametros establecidos.\n",
        estilos["nota"],
    )

    hoja.row(1).height = 1800
    hoja.row(3).height = 1040
    hoja.row(5).height = 740

    anchos = [13, 11, 13, 19, 22, 21, 28, 16, 16, 16, 22, 18, 22, 16, 26, 26]
    for col, ancho in enumerate(anchos):
        hoja.col(col).width = ancho * 256

    for col, titulo in enumerate(ENCABEZADOS):
        hoja.write(4, col, titulo, estilos["cabecera"])
        hoja.write(5, col, DESCRIPCIONES[col], estilos["descripcion"])
        estilo_requerido = "obligatorio" if REQUERIDOS[col] == "Obligatorio(*)" else "opcional"
        hoja.write(6, col, REQUERIDOS[col], estilos[estilo_requerido])

    fila = 7
    for bien in bienes:
        numero_pecosa = bien.pecosa.numero if bien.pecosa else ""
        valores = [
            bien.codigo_qr or "",
            anio,
            ejecutora,
            bien.centro_costo.ipress if bien.centro_costo else "",
            bien.persona.dni if bien.persona else "",
            bien.codigo_patrimonial,
            bien.descripcion,
            bien.fecha_alta.strftime("%Y-%m-%d") if bien.fecha_alta else "",
            bien.modelo or "",
            bien.marca or "",
            bien.estado_conservacion or "",
            bien.nro_serie or "",
            numero_pecosa,
            "",  # Color - no se usa actualmente
            "",  # Observacion Analista - no se usa actualmente
            "",  # Caracteristicas - no se usa actualmente
        ]
        for col, valor in enumerate(valores):
            hoja.write(fila, col, valor, estilos["dato"])
        fila += 1

    hoja.set_portrait(False)
    hoja.fit_width_to_pages = 1
    hoja.set_horz_split_pos(7)
    hoja.set_panes_frozen(True)
    libro.save(ruta_salida)
    return ruta_salida


def corregir_codigo_patrimonial(valor) -> str:
    """El reporte de One Visión trae el código patrimonial con un
    carácter vacío al inicio. Nos quedamos con los últimos 12 dígitos,
    que son el código real."""
    texto = str(valor).strip()
    solo_digitos = "".join(ch for ch in texto if ch.isdigit())
    return solo_digitos[-12:] if len(solo_digitos) >= 12 else solo_digitos


def _indice_columna(encabezados: list[str], candidatos: list[str]) -> int | None:
    for candidato in candidatos:
        if candidato in encabezados:
            return encabezados.index(candidato)
    return None


def _texto_excel(valor) -> str:
    """Convierte una celda a texto sin producir valores como ``None`` o ``123.0``."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def iterar_reporte_qr_onevision(ruta_archivo: str) -> Iterator[dict[str, str]]:
    """Lee progresivamente solo las columnas necesarias del reporte QR.

    ``read_only`` evita construir en memoria un DataFrame con todo el reporte.
    Esto es importante porque el archivo de One Visión puede contener muchos
    más bienes que los pertenecientes al lote que se está cruzando.
    """
    libro = load_workbook(ruta_archivo, read_only=True, data_only=True)
    try:
        hoja = libro.active
        filas = hoja.iter_rows(values_only=True)
        primera_fila = next(filas, None)
        if primera_fila is None:
            raise ValueError("El reporte de One Visión está vacío.")

        encabezados = [_texto_excel(valor) for valor in primera_fila]
        indice_codigo = _indice_columna(
            encabezados,
            ["Código Patrimonial", "Codigo Patrimonial", "codigo_patrimonial"],
        )
        if indice_codigo is None:
            raise ValueError(
                "No se encontró la columna de Código Patrimonial en el reporte de One Visión."
            )

        indice_qr = _indice_columna(encabezados, ["Código QR", "Codigo QR"])
        if indice_qr is None:
            raise ValueError(
                "No se encontró la columna de Código QR en el reporte de One Visión."
            )
        indice_ruta = _indice_columna(encabezados, ["Ruta QR", "URL", "Ruta"])

        for fila in filas:
            valor_codigo = fila[indice_codigo] if indice_codigo < len(fila) else None
            codigo = corregir_codigo_patrimonial(valor_codigo)
            if not codigo:
                continue

            valor_qr = fila[indice_qr] if indice_qr < len(fila) else None
            valor_ruta = (
                fila[indice_ruta]
                if indice_ruta is not None and indice_ruta < len(fila)
                else None
            )
            yield {
                "codigo_patrimonial_corregido": codigo,
                "codigo_qr": _texto_excel(valor_qr),
                "ruta_qr": _texto_excel(valor_ruta),
            }
    finally:
        libro.close()


def leer_reporte_qr_onevision(ruta_archivo: str):
    """Compatibilidad para usos que todavía necesiten un DataFrame completo.

    El cruce web utiliza :func:`iterar_reporte_qr_onevision` para no cargar el
    reporte entero en memoria.
    """
    import pandas as pd

    df = pd.read_excel(ruta_archivo)
    columna_codigo = None
    for candidata in ["Código Patrimonial", "Codigo Patrimonial", "codigo_patrimonial"]:
        if candidata in df.columns:
            columna_codigo = candidata
            break
    if columna_codigo is None:
        raise ValueError(
            "No se encontró la columna de Código Patrimonial en el reporte de One Visión."
        )

    df["codigo_patrimonial_corregido"] = df[columna_codigo].apply(
        corregir_codigo_patrimonial
    )
    return df
