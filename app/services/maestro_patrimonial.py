"""Validación, comparación y trazabilidad del reporte Maestro Patrimonial SIGA."""
from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database import SessionLocal
from app.models import (
    BienCargaPatrimonial,
    BienPatrimonial,
    CambioBienPatrimonial,
    CargaPatrimonial,
    ConflictoBienPatrimonial,
    CorreccionBienPatrimonial,
    VersionBienPatrimonial,
)


ENCABEZADOS_SIGA = [
    "codigo_patrimonial", "descripcion", "nombre_sede", "nombre_depend",
    "responsable", "usuario", "nombre_prov", "fecha_compra", "valor_compra",
    "fecha_alta", "valor_inicial", "sede", "pliego", "ubicac_fisica",
    "nombre_item", "sec_ejec", "tipo_modalidad", "codigo_barra", "modelo",
    "nro_orden", "medidas", "hvalor_neto", "abrev_movimto", "secuencia",
    "nro_documento", "flag_compartido", "nombre", "centro_costo", "nombre",
    "abreviatura", "fecha_nea", "tipo_doc_refer", "sec_modelo", "nro_serie",
    "grupo_bien", "clase_bien", "familia_bien", "item_bien", "color",
    "caracteristicas", "observaciones",
]

# El número es el índice de la columna en el reporte SIGA, comenzando en cero.
CAMPOS = {
    "codigo_patrimonial": (0, "texto"),
    "descripcion": (1, "texto"),
    "nombre_dependencia": (3, "texto"),
    "usuario": (5, "texto"),
    "fecha_compra": (7, "fecha"),
    "valor_compra": (8, "decimal"),
    "fecha_alta": (9, "fecha"),
    "valor_inicial": (10, "decimal"),
    "ubicacion_fisica": (13, "texto"),
    "codigo_qr": (17, "texto"),
    "modelo": (18, "texto"),
    "numero_orden": (19, "texto"),
    "medidas": (20, "texto"),
    "valor_neto": (21, "decimal"),
    "numero_documento": (24, "texto"),
    "marca": (26, "texto"),
    "estado_conservacion": (28, "texto"),
    "fecha_nea": (30, "fecha"),
    "numero_serie": (33, "texto"),
    "color": (38, "texto"),
    "caracteristicas": (39, "texto"),
    "observaciones": (40, "texto"),
}

ETIQUETAS_CAMPOS = {
    "codigo_patrimonial": "Código patrimonial",
    "codigo_qr": "Código QR",
    "descripcion": "Descripción",
    "nombre_dependencia": "Dependencia",
    "usuario": "Usuario",
    "fecha_compra": "Fecha de compra",
    "valor_compra": "Valor de compra",
    "fecha_alta": "Fecha de alta",
    "valor_inicial": "Valor inicial",
    "ubicacion_fisica": "Ubicación física",
    "modelo": "Modelo",
    "numero_orden": "Número de orden",
    "medidas": "Medidas",
    "valor_neto": "Valor neto",
    "numero_documento": "Número de documento",
    "marca": "Marca",
    "estado_conservacion": "Estado de conservación",
    "fecha_nea": "Fecha NEA",
    "numero_serie": "Número de serie",
    "color": "Color",
    "caracteristicas": "Características",
    "observaciones": "Observaciones",
}

CAMPOS_OPCIONALES = {
    "fecha_compra", "fecha_alta", "codigo_qr", "modelo", "medidas",
    "numero_documento", "numero_orden", "fecha_nea", "numero_serie", "color",
    "caracteristicas", "observaciones",
}
CAMPOS_EDITABLES = tuple(campo for campo in CAMPOS if campo != "codigo_patrimonial")
CAMPOS_COMPARABLES = tuple(CAMPOS)
TAMANO_LOTE = 750
MAX_ERRORES_DETALLE = 250


def _normalizar_encabezado(valor) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"\s+", "_", texto.strip().casefold())
    return texto


def texto_excel(valor) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    texto = " ".join(str(valor).replace("\r", " ").replace("\n", " ").split())
    return texto or None


def normalizar_codigo_qr(valor) -> str | None:
    """Quita ceros iniciales solo cuando el QR contiene únicamente dígitos."""
    texto = texto_excel(valor)
    if texto and texto.isdigit():
        return texto.lstrip("0") or "0"
    return texto


def fecha_excel(valor) -> date | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    if not texto:
        return None
    texto = texto.split(" ")[0]
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ValueError("fecha no reconocida")


def decimal_excel(valor) -> Decimal | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, Decimal):
        return valor
    if isinstance(valor, (int, float)):
        return Decimal(str(valor))
    texto = str(valor).strip().replace(" ", "")
    if not texto:
        return None
    if "," in texto and "." in texto:
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        texto = texto.replace(",", ".")
    try:
        return Decimal(texto)
    except InvalidOperation as exc:
        raise ValueError("número no reconocido") from exc


def serializar_valor(valor):
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()
    if isinstance(valor, Decimal):
        normalizado = format(valor.normalize(), "f")
        return "0" if normalizado in ("-0", "") else normalizado
    return valor


def deserializar_valor(campo: str, valor):
    tipo = CAMPOS[campo][1]
    if valor in (None, ""):
        return None
    if tipo == "fecha":
        return date.fromisoformat(str(valor))
    if tipo == "decimal":
        return Decimal(str(valor))
    return str(valor)


def valor_visible(valor) -> str | None:
    serializado = serializar_valor(valor)
    return None if serializado is None else str(serializado)


def _json_seguro(valor):
    if isinstance(valor, (date, datetime, Decimal)):
        return serializar_valor(valor)
    if isinstance(valor, (int, float, str, bool)) or valor is None:
        return valor
    return str(valor)


def calcular_huella(ruta_archivo: str) -> str:
    digest = hashlib.sha256()
    with open(ruta_archivo, "rb") as archivo:
        for bloque in iter(lambda: archivo.read(1024 * 1024), b""):
            digest.update(bloque)
    return digest.hexdigest()


def _agregar_error(errores: list[dict], fila: int | None, campo: str, mensaje: str):
    if len(errores) < MAX_ERRORES_DETALLE:
        errores.append({"fila": fila, "campo": campo, "mensaje": mensaje})


def _convertir_fila(fila, numero_fila: int):
    valores_fuente = [_json_seguro(valor) for valor in list(fila[:len(ENCABEZADOS_SIGA)])]
    if not any(valor not in (None, "") for valor in valores_fuente):
        return None, valores_fuente, []

    errores = []
    datos = {}
    for campo, (indice, tipo) in CAMPOS.items():
        valor = fila[indice] if indice < len(fila) else None
        try:
            if tipo == "fecha":
                convertido = fecha_excel(valor)
            elif tipo == "decimal":
                convertido = decimal_excel(valor)
            else:
                convertido = texto_excel(valor)
                if campo == "codigo_qr":
                    convertido = normalizar_codigo_qr(convertido)
        except ValueError as exc:
            convertido = None
            errores.append({
                "fila": numero_fila,
                "campo": ETIQUETAS_CAMPOS[campo],
                "mensaje": str(exc),
            })
        if campo not in CAMPOS_OPCIONALES and convertido is None:
            errores.append({
                "fila": numero_fila,
                "campo": ETIQUETAS_CAMPOS[campo],
                "mensaje": "Campo obligatorio vacío.",
            })
        datos[campo] = convertido
    return datos, valores_fuente, errores


def _comparables(datos: dict) -> dict:
    return {campo: serializar_valor(datos.get(campo)) for campo in CAMPOS_COMPARABLES}


def diferencias(anterior: dict, nuevo: dict) -> list[dict]:
    return [
        {
            "campo": campo,
            "valor_anterior": anterior.get(campo),
            "valor_nuevo": nuevo.get(campo),
        }
        for campo in CAMPOS_COMPARABLES
        if anterior.get(campo) != nuevo.get(campo)
    ]


def _insertar_por_lotes(db: Session, modelo, filas: list[dict]):
    for inicio in range(0, len(filas), TAMANO_LOTE):
        db.bulk_insert_mappings(modelo, filas[inicio:inicio + TAMANO_LOTE])
        db.commit()


def validar_carga_patrimonial(carga_id: int, ruta_archivo: str):
    """Valida el archivo en segundo plano y deja sus filas listas para confirmar."""
    db = SessionLocal()
    libro = None
    try:
        carga = db.get(CargaPatrimonial, carga_id)
        if carga is None:
            return
        carga.estado = "Validando"
        carga.progreso = 2
        carga.mensaje_progreso = "Abriendo y comprobando la estructura del reporte."
        db.commit()

        repetida = db.query(CargaPatrimonial.id).filter(
            CargaPatrimonial.huella_archivo == carga.huella_archivo,
            CargaPatrimonial.id != carga.id,
            CargaPatrimonial.estado.in_(("Lista para confirmar", "Procesando", "Completada")),
        ).first()
        if repetida:
            carga.estado = "Rechazada"
            carga.progreso = 100
            carga.mensaje_progreso = "El archivo ya había sido registrado."
            carga.total_errores = 1
            carga.detalle_validacion = json.dumps([{
                "fila": None,
                "campo": "Archivo",
                "mensaje": f"Este archivo ya fue registrado en la carga #{repetida[0]}.",
            }], ensure_ascii=False)
            db.commit()
            return

        libro = load_workbook(ruta_archivo, read_only=True, data_only=True)
        hoja = libro.active
        total_estimado = max(1, (hoja.max_row or 1) - 1)
        iterador = hoja.iter_rows(values_only=True)
        encabezados = next(iterador, None)
        errores: list[dict] = []
        total_errores = 0
        if encabezados is None:
            raise ValueError("El archivo está vacío.")
        recibidos = [_normalizar_encabezado(v) for v in encabezados[:len(ENCABEZADOS_SIGA)]]
        esperados = [_normalizar_encabezado(v) for v in ENCABEZADOS_SIGA]
        if recibidos != esperados:
            total_errores += 1
            _agregar_error(
                errores, 1, "Encabezados",
                "La estructura no coincide con el reporte estándar de SIGA Patrimonio.",
            )
            carga.estado = "Rechazada"
            carga.progreso = 100
            carga.mensaje_progreso = "La estructura del archivo no es válida."
            carga.total_errores = total_errores
            carga.detalle_validacion = json.dumps(errores, ensure_ascii=False)
            db.commit()
            return

        existentes = {
            codigo: json.loads(datos)
            for codigo, datos in db.query(
                BienPatrimonial.codigo_patrimonial,
                BienPatrimonial.datos_importados,
            ).all()
        }
        codigos_vistos: set[str] = set()
        qr_ocurrencias: dict[str, list[tuple[str, int]]] = defaultdict(list)
        filas_lote: list[dict] = []
        total_filas = 0
        validas = 0
        conteos = defaultdict(int)

        db.query(BienCargaPatrimonial).filter(
            BienCargaPatrimonial.carga_id == carga.id
        ).delete(synchronize_session=False)
        db.commit()

        for numero_fila, fila in enumerate(iterador, start=2):
            datos, fuente, errores_fila = _convertir_fila(fila, numero_fila)
            if datos is None:
                continue
            total_filas += 1
            if errores_fila:
                total_errores += len(errores_fila)
                for error in errores_fila:
                    _agregar_error(
                        errores, error["fila"], error["campo"], error["mensaje"]
                    )
                continue

            codigo = datos["codigo_patrimonial"]
            qr = datos.get("codigo_qr")
            if codigo in codigos_vistos:
                total_errores += 1
                _agregar_error(
                    errores, numero_fila, "Código patrimonial",
                    f"Código repetido en el archivo: {codigo}.",
                )
                continue
            codigos_vistos.add(codigo)
            if qr:
                qr_ocurrencias[qr].append((codigo, numero_fila))

            comparables = _comparables(datos)
            anterior = existentes.get(codigo)
            if anterior is None:
                clasificacion = "Nuevo"
            elif diferencias(anterior, comparables):
                clasificacion = "Actualizado"
            else:
                clasificacion = "Sin cambios"
            conteos[clasificacion] += 1
            validas += 1
            filas_lote.append({
                "carga_id": carga.id,
                "numero_fila": numero_fila,
                "codigo_patrimonial": codigo,
                "clasificacion": clasificacion,
                "datos_comparables": json.dumps(comparables, ensure_ascii=False),
                "datos_fuente": json.dumps(fuente, ensure_ascii=False),
            })
            if len(filas_lote) >= TAMANO_LOTE:
                db.bulk_insert_mappings(BienCargaPatrimonial, filas_lote)
                carga.progreso = min(95, max(5, int(total_filas * 95 / total_estimado)))
                carga.mensaje_progreso = (
                    f"Validando filas: {min(total_filas, total_estimado)} de "
                    f"{total_estimado}."
                )
                db.commit()
                filas_lote.clear()

        if filas_lote:
            db.bulk_insert_mappings(BienCargaPatrimonial, filas_lote)
            db.commit()

        alertas = []
        for qr, ocurrencias in qr_ocurrencias.items():
            if len(ocurrencias) < 2:
                continue
            codigos = ", ".join(codigo for codigo, _ in ocurrencias)
            filas_qr = ", ".join(str(fila) for _, fila in ocurrencias)
            alertas.append({
                "campo": "Código QR",
                "qr": qr,
                "codigos": [codigo for codigo, _ in ocurrencias],
                "filas": [fila for _, fila in ocurrencias],
                "mensaje": (
                    f"El QR {qr} aparece en los códigos {codigos} "
                    f"(filas {filas_qr})."
                ),
            })

        if total_errores:
            db.query(BienCargaPatrimonial).filter(
                BienCargaPatrimonial.carga_id == carga.id
            ).delete(synchronize_session=False)
            carga.estado = "Rechazada"
        else:
            no_incluidos = (
                set(existentes) - codigos_vistos
                if carga.tipo_carga == "Completa" else set()
            )
            conteos["No incluido"] = len(no_incluidos)
            filas_no_incluidas = [{
                "carga_id": carga.id,
                "numero_fila": None,
                "codigo_patrimonial": codigo,
                "clasificacion": "No incluido",
            } for codigo in sorted(no_incluidos)]
            _insertar_por_lotes(db, BienCargaPatrimonial, filas_no_incluidas)
            carga.estado = "Lista para confirmar"

        carga.total_filas = total_filas
        carga.filas_validas = validas
        carga.total_errores = total_errores
        carga.total_alertas = len(alertas)
        carga.total_nuevos = conteos["Nuevo"]
        carga.total_actualizados = conteos["Actualizado"]
        carga.total_sin_cambios = conteos["Sin cambios"]
        carga.total_no_incluidos = conteos["No incluido"]
        carga.detalle_validacion = json.dumps(errores, ensure_ascii=False)
        carga.detalle_alertas = json.dumps(alertas, ensure_ascii=False)
        carga.progreso = 100
        carga.mensaje_progreso = (
            "Validación terminada. Revisa el resultado antes de confirmar."
            if not total_errores else
            "Validación terminada con errores."
        )
        db.commit()
    except Exception as exc:
        db.rollback()
        carga = db.get(CargaPatrimonial, carga_id)
        if carga:
            db.query(BienCargaPatrimonial).filter(
                BienCargaPatrimonial.carga_id == carga.id
            ).delete(synchronize_session=False)
            carga.estado = "Rechazada"
            carga.progreso = 100
            carga.mensaje_progreso = "La validación no pudo completarse."
            carga.total_errores = max(1, carga.total_errores or 0)
            carga.detalle_validacion = json.dumps([{
                "fila": None, "campo": "Archivo", "mensaje": str(exc),
            }], ensure_ascii=False)
            db.commit()
    finally:
        if libro is not None:
            libro.close()
        if os.path.exists(ruta_archivo):
            os.remove(ruta_archivo)
        carga = db.get(CargaPatrimonial, carga_id)
        if carga is not None and carga.estado != "Validando":
            carga.archivo_contenido = None
            db.commit()
        db.close()


def _modelo_desde_comparables(datos: dict) -> dict:
    return {campo: deserializar_valor(campo, datos.get(campo)) for campo in CAMPOS}


def _consultar_ids_por_codigo(db: Session, codigos: list[str]) -> dict[str, int]:
    resultado = {}
    for inicio in range(0, len(codigos), TAMANO_LOTE):
        for codigo, bien_id in db.query(
            BienPatrimonial.codigo_patrimonial, BienPatrimonial.id
        ).filter(
            BienPatrimonial.codigo_patrimonial.in_(codigos[inicio:inicio + TAMANO_LOTE])
        ).all():
            resultado[codigo] = bien_id
    return resultado


def confirmar_carga_patrimonial(carga_id: int):
    """Aplica una carga validada sin perder correcciones ni versiones anteriores."""
    db = SessionLocal()
    try:
        carga = db.get(CargaPatrimonial, carga_id)
        if carga is None or carga.estado not in ("Lista para confirmar", "Procesando"):
            return
        carga.estado = "Procesando"
        carga.progreso = 5
        carga.mensaje_progreso = "Preparando los bienes validados."
        db.commit()

        filas = db.query(BienCargaPatrimonial).filter(
            BienCargaPatrimonial.carga_id == carga.id
        ).order_by(BienCargaPatrimonial.id).all()
        carga.progreso = 15
        carga.mensaje_progreso = f"Aplicando {len(filas)} resultados validados."
        db.commit()
        codigos_presentes = [f.codigo_patrimonial for f in filas if f.clasificacion != "No incluido"]
        existentes = {
            bien.codigo_patrimonial: bien
            for bien in db.query(BienPatrimonial).all()
        }

        nuevos = []
        for fila in filas:
            if fila.clasificacion != "Nuevo":
                continue
            datos = json.loads(fila.datos_comparables)
            nuevos.append({
                **_modelo_desde_comparables(datos),
                "datos_importados": fila.datos_comparables,
                "datos_fuente": fila.datos_fuente,
                "ultima_carga_id": carga.id,
                "creado_en": datetime.utcnow(),
                "actualizado_en": datetime.utcnow(),
            })
        for inicio in range(0, len(nuevos), TAMANO_LOTE):
            db.bulk_insert_mappings(BienPatrimonial, nuevos[inicio:inicio + TAMANO_LOTE])
        db.flush()

        ids_por_codigo = _consultar_ids_por_codigo(
            db, [fila.codigo_patrimonial for fila in filas]
        )
        ids_existentes = list(ids_por_codigo.values())
        correcciones = defaultdict(dict)
        if ids_existentes:
            for correccion in db.query(CorreccionBienPatrimonial).filter(
                CorreccionBienPatrimonial.activa == 1,
            ).all():
                correcciones[correccion.bien_id][correccion.campo] = correccion

        ahora = datetime.utcnow()
        actualizaciones = []
        versiones_pendientes = []
        cambios_por_bien: dict[int, list[dict]] = {}
        conflictos = []
        enlaces_fila = []

        for fila in filas:
            bien_id = ids_por_codigo.get(fila.codigo_patrimonial)
            if bien_id:
                enlaces_fila.append({"id": fila.id, "bien_id": bien_id})
            if fila.clasificacion == "No incluido":
                continue

            datos_nuevos = json.loads(fila.datos_comparables)
            if fila.clasificacion == "Nuevo":
                versiones_pendientes.append({
                    "bien_id": bien_id,
                    "carga_id": carga.id,
                    "origen": "Importación",
                    "usuario": carga.usuario_carga,
                    "creado_en": ahora,
                    "snapshot": fila.datos_comparables,
                })
                continue

            bien = existentes.get(fila.codigo_patrimonial)
            if bien is None:
                bien = db.get(BienPatrimonial, bien_id)
            datos_anteriores = json.loads(bien.datos_importados)
            cambios = diferencias(datos_anteriores, datos_nuevos)
            mapping = {
                "id": bien.id,
                "datos_importados": fila.datos_comparables,
                "datos_fuente": fila.datos_fuente,
                "ultima_carga_id": carga.id,
                "actualizado_en": ahora,
            }
            activas = correcciones.get(bien.id, {})
            for campo in CAMPOS:
                correccion = activas.get(campo)
                valor_siga_nuevo = datos_nuevos.get(campo)
                valor_siga_anterior = datos_anteriores.get(campo)
                if correccion is None:
                    mapping[campo] = deserializar_valor(campo, valor_siga_nuevo)
                elif valor_siga_nuevo == correccion.valor_nuevo:
                    correccion.activa = 0
                    correccion.cerrada_en = ahora
                    mapping[campo] = deserializar_valor(campo, valor_siga_nuevo)
                elif valor_siga_nuevo != valor_siga_anterior:
                    conflictos.append(ConflictoBienPatrimonial(
                        bien_id=bien.id,
                        carga_id=carga.id,
                        correccion_id=correccion.id,
                        campo=campo,
                        valor_siga_anterior=valor_visible(valor_siga_anterior),
                        valor_siga_nuevo=valor_visible(valor_siga_nuevo),
                        valor_manual=correccion.valor_nuevo,
                    ))
            actualizaciones.append(mapping)
            if fila.clasificacion == "Actualizado":
                snapshot = dict(datos_nuevos)
                for campo, correccion in activas.items():
                    if correccion.activa:
                        snapshot[campo] = correccion.valor_nuevo
                versiones_pendientes.append({
                    "bien_id": bien.id,
                    "carga_id": carga.id,
                    "origen": "Importación",
                    "usuario": carga.usuario_carga,
                    "creado_en": ahora,
                    "snapshot": json.dumps(snapshot, ensure_ascii=False),
                })
                cambios_por_bien[bien.id] = cambios

        for inicio in range(0, len(actualizaciones), TAMANO_LOTE):
            db.bulk_update_mappings(BienPatrimonial, actualizaciones[inicio:inicio + TAMANO_LOTE])
        for inicio in range(0, len(enlaces_fila), TAMANO_LOTE):
            db.bulk_update_mappings(BienCargaPatrimonial, enlaces_fila[inicio:inicio + TAMANO_LOTE])
        if conflictos:
            db.add_all(conflictos)
        for inicio in range(0, len(versiones_pendientes), TAMANO_LOTE):
            db.bulk_insert_mappings(
                VersionBienPatrimonial,
                versiones_pendientes[inicio:inicio + TAMANO_LOTE],
            )
        db.flush()

        versiones_ids = {
            bien_id: version_id
            for bien_id, version_id in db.query(
                VersionBienPatrimonial.bien_id, VersionBienPatrimonial.id
            ).filter(VersionBienPatrimonial.carga_id == carga.id).all()
        }
        cambios_insertar = []
        for bien_id, cambios in cambios_por_bien.items():
            version_id = versiones_ids.get(bien_id)
            if not version_id:
                continue
            for cambio in cambios:
                cambios_insertar.append({
                    "version_id": version_id,
                    "campo": cambio["campo"],
                    "valor_anterior": valor_visible(cambio["valor_anterior"]),
                    "valor_nuevo": valor_visible(cambio["valor_nuevo"]),
                })
        for inicio in range(0, len(cambios_insertar), TAMANO_LOTE):
            db.bulk_insert_mappings(
                CambioBienPatrimonial,
                cambios_insertar[inicio:inicio + TAMANO_LOTE],
            )

        carga.estado = "Completada"
        carga.confirmado_en = ahora
        carga.progreso = 100
        carga.mensaje_progreso = "Carga aplicada correctamente."
        db.commit()
    except Exception as exc:
        db.rollback()
        carga = db.get(CargaPatrimonial, carga_id)
        if carga:
            carga.estado = "Interrumpida"
            carga.progreso = 0
            carga.mensaje_progreso = "La carga se reintentará desde sus datos validados."
            detalle = json.loads(carga.detalle_validacion or "[]")
            detalle.append({
                "fila": None,
                "campo": "Proceso",
                "mensaje": (
                    "No se pudo aplicar la carga. El proceso fue revertido sin "
                    f"modificar los bienes. Tipo de error: {type(exc).__name__}."
                ),
            })
            carga.detalle_validacion = json.dumps(detalle[-MAX_ERRORES_DETALLE:], ensure_ascii=False)
            db.commit()
    finally:
        db.close()


def snapshot_vigente(bien: BienPatrimonial) -> dict:
    return {campo: serializar_valor(getattr(bien, campo)) for campo in CAMPOS}


def convertir_edicion(campo: str, valor):
    tipo = CAMPOS[campo][1]
    if tipo == "fecha":
        convertido = fecha_excel(valor)
    elif tipo == "decimal":
        convertido = decimal_excel(valor)
    else:
        convertido = texto_excel(valor)
        if campo == "codigo_qr":
            convertido = normalizar_codigo_qr(convertido)
    if campo not in CAMPOS_OPCIONALES and convertido is None:
        raise ValueError(f"{ETIQUETAS_CAMPOS[campo]} es obligatorio.")
    return convertido


def editar_bien_patrimonial(
    db: Session,
    bien: BienPatrimonial,
    valores: dict,
    motivo: str,
    usuario: str,
) -> int:
    motivo = (motivo or "").strip()
    if not motivo:
        raise ValueError("El motivo de la corrección es obligatorio.")
    convertidos = {
        campo: convertir_edicion(campo, valores.get(campo))
        for campo in CAMPOS_EDITABLES
    }
    qr = convertidos.get("codigo_qr")
    if qr:
        repetido = db.query(BienPatrimonial).filter(
            BienPatrimonial.codigo_qr == qr,
            BienPatrimonial.id != bien.id,
        ).first()
        if repetido:
            raise ValueError(
                f"El QR {qr} ya está asociado al código patrimonial "
                f"{repetido.codigo_patrimonial}."
            )

    cambios = []
    ahora = datetime.utcnow()
    for campo, nuevo in convertidos.items():
        anterior = getattr(bien, campo)
        anterior_serializado = valor_visible(anterior)
        nuevo_serializado = valor_visible(nuevo)
        if anterior_serializado == nuevo_serializado:
            continue
        db.query(CorreccionBienPatrimonial).filter(
            CorreccionBienPatrimonial.bien_id == bien.id,
            CorreccionBienPatrimonial.campo == campo,
            CorreccionBienPatrimonial.activa == 1,
        ).update({
            CorreccionBienPatrimonial.activa: 0,
            CorreccionBienPatrimonial.cerrada_en: ahora,
        }, synchronize_session=False)
        db.add(CorreccionBienPatrimonial(
            bien_id=bien.id,
            campo=campo,
            valor_anterior=anterior_serializado,
            valor_nuevo=nuevo_serializado,
            motivo=motivo,
            usuario=usuario,
        ))
        setattr(bien, campo, nuevo)
        cambios.append((campo, anterior_serializado, nuevo_serializado))

    if not cambios:
        return 0
    bien.actualizado_en = ahora
    db.flush()
    version = VersionBienPatrimonial(
        bien_id=bien.id,
        carga_id=None,
        origen="Edición manual",
        usuario=usuario,
        creado_en=ahora,
        snapshot=json.dumps(snapshot_vigente(bien), ensure_ascii=False),
    )
    db.add(version)
    db.flush()
    db.add_all([
        CambioBienPatrimonial(
            version_id=version.id,
            campo=campo,
            valor_anterior=anterior,
            valor_nuevo=nuevo,
        )
        for campo, anterior, nuevo in cambios
    ])
    db.commit()
    return len(cambios)


def resolver_conflicto(
    db: Session,
    conflicto: ConflictoBienPatrimonial,
    decision: str,
    usuario: str,
):
    if conflicto.estado != "Pendiente":
        raise ValueError("El conflicto ya fue resuelto.")
    if decision not in ("manual", "siga"):
        raise ValueError("Decisión de conflicto no válida.")
    ahora = datetime.utcnow()
    conflicto.resuelto_por = usuario
    conflicto.resuelto_en = ahora
    if decision == "manual":
        conflicto.estado = "Corrección conservada"
        db.commit()
        return

    bien = conflicto.bien
    campo = conflicto.campo
    anterior = valor_visible(getattr(bien, campo))
    nuevo = conflicto.valor_siga_nuevo
    setattr(bien, campo, deserializar_valor(campo, nuevo))
    bien.actualizado_en = ahora
    conflicto.estado = "SIGA aceptado"
    conflicto.correccion.activa = 0
    conflicto.correccion.cerrada_en = ahora
    version = VersionBienPatrimonial(
        bien_id=bien.id,
        origen="Resolución de conflicto",
        usuario=usuario,
        creado_en=ahora,
        snapshot=json.dumps(snapshot_vigente(bien), ensure_ascii=False),
    )
    db.add(version)
    db.flush()
    db.add(CambioBienPatrimonial(
        version_id=version.id,
        campo=campo,
        valor_anterior=anterior,
        valor_nuevo=nuevo,
    ))
    db.commit()


def preparar_carga(
    db: Session,
    ruta: str,
    nombre: str,
    usuario: str,
    tipo_carga: str = "Completa",
) -> CargaPatrimonial:
    if tipo_carga not in ("Completa", "Parcial"):
        raise ValueError("El tipo de carga no es válido.")
    carga = CargaPatrimonial(
        nombre_archivo=Path(nombre).name,
        huella_archivo=calcular_huella(ruta),
        usuario_carga=usuario,
        estado="Validando",
        tipo_carga=tipo_carga,
        progreso=0,
        mensaje_progreso="Pendiente de validación.",
        archivo_contenido=Path(ruta).read_bytes(),
    )
    db.add(carga)
    db.commit()
    db.refresh(carga)
    return carga


def validar_carga_patrimonial_desde_bd(carga_id: int):
    """Reconstruye el temporal desde la base para iniciar o reanudar la validación."""
    db = SessionLocal()
    ruta = ""
    try:
        carga = db.get(CargaPatrimonial, carga_id)
        if carga is None or carga.estado != "Validando":
            return
        if not carga.archivo_contenido:
            carga.estado = "Interrumpida"
            carga.progreso = 0
            carga.mensaje_progreso = "El archivo temporal no está disponible para reanudar."
            db.commit()
            return
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as temporal:
            temporal.write(carga.archivo_contenido)
            ruta = temporal.name
    finally:
        db.close()
    if ruta:
        validar_carga_patrimonial(carga_id, ruta)
