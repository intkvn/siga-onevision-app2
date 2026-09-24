import io
import os
import inspect
import json
import re
import tempfile
import unittest
from datetime import date, datetime
from types import SimpleNamespace
from unittest.mock import patch

import xlrd
import xlwt
import pandas as pd
from openpyxl import Workbook as OpenpyxlWorkbook, load_workbook
from reportlab.lib.units import mm
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base, _opciones_engine
from app.models import (
    BienAlta, CentroCosto, CorreccionAsignacionBien, Expediente, Pecosa,
    LoteCarga, Persona, RelacionPecosaItem, VerificacionPecosaSiga,
    ObservacionControlPecosa, PerfilImpresionEtiqueta, InventarioImpresion,
    BienInventarioImpresion, LoteImpresionInventario,
    ItemLoteImpresionInventario, CargaPatrimonial, BienPatrimonial,
    BienCargaPatrimonial, VersionBienPatrimonial, CambioBienPatrimonial,
    CorreccionBienPatrimonial, ConflictoBienPatrimonial,
    ExportacionPatrimonial,
)
from app.routers.carga_inicial import _estado_maestros
from app.routers.control import (
    ESTADO_COMPLETA, ESTADO_EXCESO, ESTADO_FALTA_FIRMA, ESTADO_OBSERVADA,
    ESTADO_PENDIENTE_ALMACEN,
    _calcular_control, _filtrar_filas_control, _mover_bien_a_pecosa,
)
from app.routers.normalizacion import (
    _completar_bienes_lote, _diferir_pecosa_faltante, _indicadores_cruce_lote,
    _pecosas_no_encontradas,
    _regularizar_bienes_historicos, _resumen_lote,
)
from app.routers.impresion import _validar_valores_perfil, procesar_reporte_qr
from app.routers.pecosas import registrar_pecosas_multiples
from app.services.excel_relacion_pecosas import COLUMNAS_NECESARIAS, leer_relacion_pecosas
from app.services.excel_onevision import (
    ENCABEZADOS, generar_formato_importacion, iterar_reporte_qr_onevision,
)
from app.services.excel_verificacion import leer_reporte_verificacion
from app.services.excel_inventario_impresion import importar_reporte_inventario
from app.services.lote_status import expedientes_de_lotes
from app.services.pagination import paginas_visibles, rango_registros
from app.services.pdf_etiquetas import (
    _ajustar_bloques_inferiores, clasificar_bienes_impresion, generar_pdf_etiquetas,
    numero_paginas_para_bienes,
)
from app.services.pdf_ficha_patrimonial import generar_ficha_activo_pdf
from app.routers.verificacion import (
    ESTADO_CORRECTA, ESTADO_INCORRECTA, _filas_verificacion,
)
from app.routers.control_impresion import (
    _confirmar_items_impresos, _estado_bien, _marcar_pdf_generado,
    _opciones_distintas,
)
from app.routers.maestro_patrimonial import (
    _aplicar_filtros, _datos_firmante, _resumen_calidad_datos,
    generar_exportacion_patrimonial,
)
from app.services.maestro_patrimonial import (
    CAMPOS_EDITABLES, ENCABEZADOS_SIGA, confirmar_carga_patrimonial,
    editar_bien_patrimonial, preparar_carga, resolver_conflicto,
    validar_carga_patrimonial, validar_carga_patrimonial_desde_bd,
)


class ValidacionCargaInicialTest(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()

    def tearDown(self):
        self.db.close()

    def test_exige_personas_y_centros_de_costo(self):
        self.assertFalse(_estado_maestros(self.db)["maestros_listos"])

        self.db.add(Persona(nombre_completo="PERSONA DE PRUEBA", dni="12345678"))
        self.db.commit()
        self.assertFalse(_estado_maestros(self.db)["maestros_listos"])

        self.db.add(CentroCosto(nombre_depend="CENTRO DE PRUEBA", ipress="001"))
        self.db.commit()
        self.assertTrue(_estado_maestros(self.db)["maestros_listos"])


class ConfiguracionConexionTest(unittest.TestCase):
    def test_postgresql_verifica_y_recicla_conexiones(self):
        opciones = _opciones_engine("postgresql://usuario:clave@servidor/base")

        self.assertTrue(opciones["pool_pre_ping"])
        self.assertEqual(opciones["pool_recycle"], 240)

    def test_sqlite_conserva_su_configuracion_local(self):
        opciones = _opciones_engine("sqlite:///./local_dev.db")

        self.assertEqual(
            opciones,
            {"connect_args": {"check_same_thread": False}},
        )


class LecturaReporteQrTest(unittest.TestCase):
    def _guardar_libro(self, encabezados, filas):
        descriptor, ruta = tempfile.mkstemp(suffix=".xlsx")
        os.close(descriptor)
        libro = OpenpyxlWorkbook()
        hoja = libro.active
        hoja.append(encabezados)
        for fila in filas:
            hoja.append(fila)
        libro.save(ruta)
        libro.close()
        self.addCleanup(lambda: os.path.exists(ruta) and os.remove(ruta))
        return ruta

    def test_lee_progresivamente_solo_las_columnas_del_cruce(self):
        ruta = self._guardar_libro(
            ["Código Patrimonial", "Descripción", "Código QR", "Ruta QR"],
            [
                [" 740899502020", "BIEN 1", 696554, "https://qr/696554"],
                ["740899502021", "BIEN 2", None, None],
            ],
        )

        filas = list(iterar_reporte_qr_onevision(ruta))

        self.assertEqual(filas, [
            {
                "codigo_patrimonial_corregido": "740899502020",
                "codigo_qr": "696554",
                "ruta_qr": "https://qr/696554",
            },
            {
                "codigo_patrimonial_corregido": "740899502021",
                "codigo_qr": "",
                "ruta_qr": "",
            },
        ])

    def test_rechaza_reporte_sin_codigo_patrimonial(self):
        ruta = self._guardar_libro(
            ["Descripción", "Código QR"],
            [["BIEN 1", "696554"]],
        )

        with self.assertRaisesRegex(ValueError, "Código Patrimonial"):
            list(iterar_reporte_qr_onevision(ruta))

    def test_rechaza_reporte_sin_codigo_qr(self):
        ruta = self._guardar_libro(
            ["Código Patrimonial", "Descripción"],
            [["740899502020", "BIEN 1"]],
        )

        with self.assertRaisesRegex(ValueError, "Código QR"):
            list(iterar_reporte_qr_onevision(ruta))

    def test_el_cruce_no_bloquea_el_event_loop(self):
        self.assertFalse(inspect.iscoroutinefunction(procesar_reporte_qr))


class PdfEtiquetasTest(unittest.TestCase):
    def setUp(self):
        self.perfil = PerfilImpresionEtiqueta(
            nombre="Argox iX4-250 203 dpi",
            ancho_pagina_mm=105.1,
            ancho_etiqueta_mm=50.8,
            alto_etiqueta_mm=38.1,
            margen_izquierdo_mm=0.75,
            margen_derecho_mm=0.75,
            separacion_central_mm=2.0,
            avance_adicional_mm=0.0,
            desplazamiento_x_mm=0.0,
            desplazamiento_y_mm=0.0,
            rotacion_contenido=180,
            anio_1="2026",
            anio_2="2027",
            anio_marcado="2026",
        )

    def _bien(self, indice=1, ruta=None):
        return SimpleNamespace(
            id=indice,
            codigo_qr=f"69735{indice}",
            codigo_patrimonial=f"74648187373{indice}",
            ruta_qr=ruta or f"https://sir.example/qr/equipo/69735{indice}",
            descripcion="UNIDAD CENTRAL DE PROCESO CON MEMORIA Y ALMACENAMIENTO",
            centro_costo=SimpleNamespace(
                nombre_depend="LABORATORIO DE REFERENCIA REGIONAL DE SALUD PUBLICA"
            ),
        )

    def test_genera_una_pagina_por_cada_pareja(self):
        contenido = generar_pdf_etiquetas(
            [self._bien(1), self._bien(2), self._bien(3)], self.perfil,
        )

        self.assertTrue(contenido.startswith(b"%PDF-"))
        paginas = re.findall(rb"/Type\s*/Page\b", contenido)
        self.assertEqual(len(paginas), 2)
        self.assertEqual(numero_paginas_para_bienes(3), 2)
        caja = re.search(
            rb"/MediaBox\s*\[\s*0\s+0\s+([0-9.]+)\s+([0-9.]+)\s*\]",
            contenido,
        )
        self.assertIsNotNone(caja)
        ancho_mm = float(caja.group(1)) * 25.4 / 72
        alto_mm = float(caja.group(2)) * 25.4 / 72
        self.assertAlmostEqual(ancho_mm, 38.1, places=2)
        self.assertAlmostEqual(alto_mm, 105.1, places=2)
        self.assertRegex(contenido, rb"/Rotate\s+90\b")

    def test_entrega_al_qr_la_ruta_original_sin_modificar(self):
        ruta = "https://sir.example/QR/Equipo/697351?origen=OneVision"
        with patch("app.services.pdf_etiquetas._dibujar_qr") as dibujar_qr:
            generar_pdf_etiquetas([self._bien(1, ruta=ruta)], self.perfil)

        self.assertEqual(dibujar_qr.call_args.args[1], ruta)

    def test_qr_usa_el_tamano_ampliado(self):
        with patch("app.services.pdf_etiquetas._dibujar_qr") as dibujar_qr:
            generar_pdf_etiquetas([self._bien(1)], self.perfil)

        tamano_qr = dibujar_qr.call_args.args[4]
        self.assertAlmostEqual(tamano_qr / mm, 27.0, places=2)

    def test_descripcion_y_establecimiento_comparten_espacio_sin_solaparse(self):
        alto_disponible = 29.0
        ajuste = _ajustar_bloques_inferiores(
            "UNIDAD CENTRAL DE PROCESO CON MEMORIA Y ALMACENAMIENTO",
            "LABORATORIO DE REFERENCIA REGIONAL DE SALUD PUBLICA",
            35.1 * mm,
            alto_disponible,
        )

        self.assertIsNotNone(ajuste)
        descripcion, establecimiento, separacion = ajuste
        alto_usado = (
            len(descripcion[0]) * descripcion[2]
            + separacion
            + len(establecimiento[0]) * establecimiento[2]
        )
        self.assertLessEqual(alto_usado, alto_disponible)
        self.assertGreaterEqual(len(descripcion[0]), 2)
        self.assertGreaterEqual(len(establecimiento[0]), 2)

    def test_dibuja_las_dos_posiciones_de_la_pagina(self):
        with patch(
            "app.services.pdf_etiquetas._dibujar_etiqueta_logica"
        ) as dibujar:
            generar_pdf_etiquetas(
                [self._bien(1), self._bien(2)], self.perfil,
            )

        self.assertEqual(dibujar.call_count, 2)
        posiciones_y = [llamada.args[3] for llamada in dibujar.call_args_list]
        self.assertEqual(posiciones_y, [0.75, 53.55])

    def test_excluye_rutas_vacias_o_con_espacios_exteriores(self):
        bienes = [
            self._bien(1),
            self._bien(2, ruta=""),
            self._bien(3, ruta=" https://sir.example/qr/3 "),
        ]
        bienes[1].ruta_qr = ""

        imprimibles, excluidos = clasificar_bienes_impresion(bienes)

        self.assertEqual([bien.id for bien in imprimibles], [1])
        self.assertEqual(len(excluidos), 2)
        self.assertEqual(excluidos[0]["razon"], "Sin Ruta QR")
        self.assertIn("espacios", excluidos[1]["razon"])

    def test_valida_geometria_y_anios_del_perfil(self):
        valores = {
            "ancho_pagina_mm": 105.1,
            "ancho_etiqueta_mm": 50.8,
            "alto_etiqueta_mm": 38.1,
            "margen_izquierdo_mm": 0.75,
            "margen_derecho_mm": 0.75,
            "separacion_central_mm": 2.0,
            "avance_adicional_mm": 0.0,
            "desplazamiento_x_mm": 0.0,
            "desplazamiento_y_mm": 0.0,
            "rotacion_contenido": 180,
            "anio_1": "2026",
            "anio_2": "2027",
            "anio_marcado": "2026",
        }
        self.assertIsNone(_validar_valores_perfil(**valores))

        valores["separacion_central_mm"] = 3.0
        self.assertIn("exceden", _validar_valores_perfil(**valores))


class IndicadoresCruceLoteTest(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()

    def tearDown(self):
        self.db.close()

    def test_resume_total_y_bienes_por_expediente(self):
        expediente_1 = Expediente(numero="60144")
        expediente_2 = Expediente(numero="60150")
        persona = Persona(nombre_completo="PERSONA COMPLETA", dni="12345678")
        centro = CentroCosto(nombre_depend="CENTRO COMPLETO", ipress="10460")
        lote = LoteCarga(
            anio="2026", ejecutora="785", pecosas_solicitadas="4719,4720,4721",
        )
        self.db.add_all([expediente_1, expediente_2, persona, centro, lote])
        self.db.flush()

        pecosa_1 = Pecosa(numero="4719", expediente_id=expediente_1.id)
        pecosa_2 = Pecosa(numero="4720", expediente_id=expediente_1.id)
        pecosa_sin_bienes = Pecosa(numero="4721", expediente_id=expediente_2.id)
        self.db.add_all([pecosa_1, pecosa_2, pecosa_sin_bienes])
        self.db.flush()

        bienes = [
            BienAlta(
                pecosa_id=pecosa_1.id, lote_id=lote.id,
                codigo_patrimonial="740899502020", descripcion="BIEN COMPLETO 1",
                persona_id=persona.id, centro_costo_id=centro.id,
            ),
            BienAlta(
                pecosa_id=pecosa_1.id, lote_id=lote.id,
                codigo_patrimonial="740899502021", descripcion="BIEN PENDIENTE",
                centro_costo_id=centro.id,
            ),
            BienAlta(
                pecosa_id=pecosa_2.id, lote_id=lote.id,
                codigo_patrimonial="740899502022", descripcion="BIEN COMPLETO 2",
                persona_id=persona.id, centro_costo_id=centro.id,
            ),
        ]
        self.db.add_all(bienes)
        self.db.commit()

        resultado = _indicadores_cruce_lote(
            bienes, [pecosa_1, pecosa_2, pecosa_sin_bienes]
        )

        self.assertEqual(resultado["total"], 3)
        self.assertEqual(resultado["correctos"], 2)
        self.assertEqual(resultado["pendientes"], 1)
        self.assertEqual(resultado["por_expediente"], [
            {
                "expediente": "60144",
                "pecosas": ["4719", "4720"],
                "cantidad_pecosas": 2,
                "total": 3,
                "correctos": 2,
                "pendientes": 1,
            },
            {
                "expediente": "60150",
                "pecosas": ["4721"],
                "cantidad_pecosas": 1,
                "total": 0,
                "correctos": 0,
                "pendientes": 0,
            },
        ])


class PaginacionTest(unittest.TestCase):
    def test_muestra_todas_las_paginas_cuando_son_pocas(self):
        self.assertEqual(paginas_visibles(3, 5), [1, 2, 3, 4, 5])

    def test_compacta_paginas_intermedias_con_separadores(self):
        self.assertEqual(
            paginas_visibles(9, 17),
            [1, 2, None, 7, 8, 9, 10, 11, None, 16, 17],
        )

    def test_muestra_extremos_al_inicio_y_final(self):
        self.assertEqual(
            paginas_visibles(1, 17),
            [1, 2, 3, 4, 5, 6, 7, None, 16, 17],
        )
        self.assertEqual(
            paginas_visibles(17, 17),
            [1, 2, None, 11, 12, 13, 14, 15, 16, 17],
        )

    def test_calcula_rango_visible(self):
        self.assertEqual(rango_registros(2, 50, 123), (51, 100))
        self.assertEqual(rango_registros(3, 50, 123), (101, 123))
        self.assertEqual(rango_registros(1, 50, 0), (0, 0))


class FiltrosControlTest(unittest.TestCase):
    def setUp(self):
        self.filas = [
            {
                "nro_pecosa": "1200", "estado": ESTADO_COMPLETA,
                "lotes": [22], "expediente_alta": "31863",
                "expediente_firma": "52607",
            },
            {
                "nro_pecosa": "1212", "estado": ESTADO_FALTA_FIRMA,
                "lotes": [22], "expediente_alta": "31864",
                "expediente_firma": "52607",
            },
            {
                "nro_pecosa": "1227", "estado": ESTADO_COMPLETA,
                "lotes": [23], "expediente_alta": "31865",
                "expediente_firma": "53001",
            },
        ]

    def test_filtra_por_numero_parcial_de_expediente_firma(self):
        resultado = _filtrar_filas_control(
            self.filas, expediente_firma="526"
        )

        self.assertEqual(
            [fila["nro_pecosa"] for fila in resultado], ["1200", "1212"]
        )

    def test_combina_expediente_firma_con_estado_firmada(self):
        resultado = _filtrar_filas_control(
            self.filas,
            expediente_firma="52607",
            estado=ESTADO_COMPLETA,
        )

        self.assertEqual(
            [fila["nro_pecosa"] for fila in resultado], ["1200"]
        )


class RegistroMasivoPecosasTest(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()

    def tearDown(self):
        self.db.close()

    def test_no_crea_expediente_si_todas_las_pecosas_ya_existen(self):
        expediente_original = Expediente(numero="50001")
        self.db.add(expediente_original)
        self.db.flush()
        self.db.add(Pecosa(numero="4602", expediente_id=expediente_original.id))
        self.db.commit()

        respuesta = registrar_pecosas_multiples(
            request=SimpleNamespace(),
            numero_expediente="123",
            numeros_pecosa="4602",
            db=self.db,
            _=None,
        )

        self.assertEqual(respuesta.status_code, 303)
        self.assertIn("error=", respuesta.headers["location"])
        self.assertIsNone(
            self.db.query(Expediente).filter(Expediente.numero == "123").first()
        )
        self.assertEqual(self.db.query(Pecosa).count(), 1)

    def test_agrega_solo_nuevas_a_un_expediente_existente(self):
        expediente = Expediente(numero="50002")
        self.db.add(expediente)
        self.db.flush()
        self.db.add(Pecosa(numero="100", expediente_id=expediente.id))
        self.db.commit()

        respuesta = registrar_pecosas_multiples(
            request=SimpleNamespace(),
            numero_expediente="50002",
            numeros_pecosa="100\n101\n101,102",
            db=self.db,
            _=None,
        )

        self.assertEqual(respuesta.status_code, 303)
        pecosas = self.db.query(Pecosa).order_by(Pecosa.numero).all()
        self.assertEqual([pecosa.numero for pecosa in pecosas], ["100", "101", "102"])
        self.assertTrue(all(pecosa.expediente_id == expediente.id for pecosa in pecosas))
        self.assertEqual(
            self.db.query(Expediente).filter(Expediente.numero == "50002").count(),
            1,
        )


class CorreccionPecosaTest(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()

        expediente_origen = Expediente(numero="41724")
        expediente_destino = Expediente(numero="43613")
        self.db.add_all([expediente_origen, expediente_destino])
        self.db.flush()
        lote = LoteCarga(id=14, anio="2026", ejecutora="785", pecosas_solicitadas="2011")
        self.db.add(lote)
        self.origen = Pecosa(numero="2011", expediente_id=expediente_origen.id, estado="StickerGenerado")
        self.destino = Pecosa(numero="2473", expediente_id=expediente_destino.id)
        self.db.add_all([self.origen, self.destino])
        self.db.flush()
        self.bien_corregido = BienAlta(
            pecosa_id=self.origen.id,
            lote_id=14,
            codigo_patrimonial="536498312255",
            descripcion="TERMO PARA TRANSPORTE DE BIOLOGICOS Y VACUNAS",
            codigo_qr="695614",
        )
        self.db.add_all([
            BienAlta(
                pecosa_id=self.origen.id,
                lote_id=14,
                codigo_patrimonial="536498312210",
                descripcion="TERMO PARA TRANSPORTE DE BIOLOGICOS Y VACUNAS",
                codigo_qr="695613",
            ),
            self.bien_corregido,
            RelacionPecosaItem(nro_pecosa="2011", cant_aprobada=1),
        ])
        self.db.commit()

    def tearDown(self):
        self.db.close()

    def test_detecta_exceso_y_guarda_historial_al_reasignar(self):
        control_inicial = {fila["nro_pecosa"]: fila for fila in _calcular_control(self.db)}
        self.assertEqual(control_inicial["2011"]["estado"], ESTADO_EXCESO)
        self.assertEqual(control_inicial["2011"]["cantidad_ingresada"], 2)
        self.assertEqual(control_inicial["2011"]["lotes"], [14])

        _mover_bien_a_pecosa(
            self.db, self.bien_corregido, self.destino,
            "SIGA MP corrigió la pecosa asignada al bien.",
        )
        self.db.commit()

        self.assertEqual(self.bien_corregido.pecosa_id, self.destino.id)
        self.assertEqual(self.destino.estado, "StickerGenerado")
        historial = self.db.query(CorreccionAsignacionBien).one()
        self.assertEqual(historial.pecosa_origen_id, self.origen.id)
        self.assertEqual(historial.pecosa_destino_id, self.destino.id)
        self.assertEqual(historial.motivo, "SIGA MP corrigió la pecosa asignada al bien.")
        self.assertIn("2473", self.db.query(LoteCarga).get(14).pecosas_solicitadas.split(","))

        control_final = {fila["nro_pecosa"]: fila for fila in _calcular_control(self.db)}
        self.assertEqual(control_final["2011"]["cantidad_ingresada"], 1)
        self.assertNotEqual(control_final["2011"]["estado"], ESTADO_EXCESO)


class LecturaRelacionPecosasXlsTest(unittest.TestCase):
    @patch("app.services.excel_relacion_pecosas.pd.read_excel")
    def test_selecciona_xlrd_explicitamente_para_xls(self, leer_excel):
        leer_excel.return_value.columns = COLUMNAS_NECESARIAS

        leer_relacion_pecosas("reporte_siga.xls")

        leer_excel.assert_called_once_with("reporte_siga.xls", engine="xlrd")

    def test_lee_archivo_xls_de_siga(self):
        libro = xlwt.Workbook()
        hoja = libro.add_sheet("Relación")
        for columna, encabezado in enumerate(COLUMNAS_NECESARIAS):
            hoja.write(0, columna, encabezado)
            hoja.write(1, columna, 1 if encabezado == "cant_aprobada" else "DATO")

        descriptor, ruta = tempfile.mkstemp(suffix=".xls")
        os.close(descriptor)
        try:
            libro.save(ruta)
            resultado = leer_relacion_pecosas(ruta)
        finally:
            os.remove(ruta)

        self.assertEqual(list(resultado.columns), COLUMNAS_NECESARIAS)
        self.assertEqual(len(resultado), 1)
        self.assertEqual(resultado.iloc[0]["cant_aprobada"], 1)


class FormatoImportacionOneVisionTest(unittest.TestCase):
    def test_generar_archivo_con_qr_y_formato_vigente(self):
        bien = SimpleNamespace(
            codigo_qr="QR-001",
            codigo_patrimonial="740899502020",
            descripcion="UNIDAD CENTRAL",
            fecha_alta=date(2026, 6, 22),
            modelo="M90",
            marca="LENOVO",
            estado_conservacion="Bueno",
            nro_serie="SN001",
            pecosa=SimpleNamespace(numero="1417"),
            centro_costo=SimpleNamespace(ipress="4451"),
            persona=SimpleNamespace(dni="44797419"),
        )
        descriptor, ruta = tempfile.mkstemp(suffix=".xls")
        os.close(descriptor)
        try:
            generar_formato_importacion([bien], "2026", "785", ruta)
            libro = xlrd.open_workbook(ruta, formatting_info=True)
            hoja = libro.sheet_by_index(0)
        finally:
            os.remove(ruta)

        self.assertEqual(hoja.ncols, 16)
        self.assertEqual(hoja.row_values(4), ENCABEZADOS)
        self.assertEqual(hoja.cell_value(7, 0), "QR-001")
        self.assertEqual(hoja.cell_value(7, 1), "2026")
        self.assertEqual(hoja.cell_value(7, 2), "785")
        self.assertEqual(hoja.cell_value(7, 5), "740899502020")
        self.assertEqual(hoja.cell_value(7, 12), "1417")
        self.assertNotEqual(
            libro.xf_list[hoja.cell_xf_index(4, 0)].background.pattern_colour_index,
            0,
        )


class VerificacionPecosasTest(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()
        expediente = Expediente(numero="50001")
        lote = LoteCarga(anio="2026", ejecutora="785", pecosas_solicitadas="2473")
        self.db.add_all([expediente, lote])
        self.db.flush()
        pecosa = Pecosa(numero="2473", expediente_id=expediente.id, estado="Firmada")
        self.db.add(pecosa)
        self.db.flush()
        self.db.add_all([
            BienAlta(
                pecosa_id=pecosa.id, lote_id=lote.id,
                codigo_patrimonial="740899502037", descripcion="BIEN DE PRUEBA",
            ),
            RelacionPecosaItem(nro_pecosa="2473", ano_eje="2026", cant_aprobada=1),
            VerificacionPecosaSiga(
                codigo_patrimonial="740899502037", nro_pecosa="2473", anio_siga="2026",
            ),
        ])
        self.db.commit()

    def tearDown(self):
        self.db.close()

    def test_verifica_pecosa_y_anio_en_columnas_independientes(self):
        fila = _filas_verificacion(self.db)[0]
        self.assertEqual(fila["estado"], ESTADO_CORRECTA)
        self.assertEqual(fila["anio_lote"], "2026")
        self.assertEqual(fila["anio_siga"], "2026")
        self.assertEqual(fila["pecosa_alta"], "2473")
        self.assertEqual(fila["pecosa_real"], "2473")

        registro = self.db.query(VerificacionPecosaSiga).one()
        registro.nro_pecosa = "2474"
        self.db.commit()
        self.assertEqual(_filas_verificacion(self.db)[0]["estado"], ESTADO_INCORRECTA)

    def test_lee_reporte_xls_con_anio_separado(self):
        libro = xlwt.Workbook()
        hoja = libro.add_sheet("SIGA")
        for columna, valor in enumerate(["codigo_patrimonial", "nro_pecosa", "fecha_alta"]):
            hoja.write(0, columna, valor)
        hoja.write(1, 0, "740899502037")
        hoja.write(1, 1, "2473")
        hoja.write(1, 2, "27/08/2026 00:00:00")
        descriptor, ruta = tempfile.mkstemp(suffix=".xls")
        os.close(descriptor)
        try:
            libro.save(ruta)
            reporte = leer_reporte_verificacion(ruta)
        finally:
            os.remove(ruta)

        self.assertEqual(reporte.iloc[0]["codigo_patrimonial"], "740899502037")
        self.assertEqual(reporte.iloc[0]["nro_pecosa"], "2473")
        self.assertEqual(reporte.iloc[0]["anio_siga"], "2026")


class ObservacionControlPecosaTest(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()
        self.db.add(RelacionPecosaItem(
            nro_pecosa="36", ano_eje="2026", cant_aprobada=1,
        ))
        self.db.commit()

    def tearDown(self):
        self.db.close()

    def test_observacion_excluye_pendiente_por_anio_y_pecosa(self):
        self.assertEqual(_calcular_control(self.db)[0]["estado"], ESTADO_PENDIENTE_ALMACEN)
        self.db.add(ObservacionControlPecosa(
            ano_eje="2026", nro_pecosa="36",
            causal="Transferencia a otra RIS/unidad ejecutora",
            sustento="La RIS maneja su propio SIGA.",
        ))
        self.db.commit()

        fila = _calcular_control(self.db)[0]
        self.assertEqual(fila["estado"], ESTADO_OBSERVADA)
        self.assertEqual(fila["observacion"].sustento, "La RIS maneja su propio SIGA.")


class ExpedientesPorLotesTest(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()

        expediente_uno = Expediente(numero="40001")
        expediente_dos = Expediente(numero="40002")
        lote_uno = LoteCarga(anio="2026", ejecutora="785", pecosas_solicitadas="101,102")
        lote_dos = LoteCarga(anio="2026", ejecutora="785", pecosas_solicitadas="102,999")
        self.db.add_all([expediente_uno, expediente_dos, lote_uno, lote_dos])
        self.db.flush()
        self.db.add_all([
            Pecosa(numero="101", expediente_id=expediente_uno.id),
            Pecosa(numero="102", expediente_id=expediente_dos.id),
        ])
        self.db.commit()
        self.lote_uno = lote_uno
        self.lote_dos = lote_dos

    def tearDown(self):
        self.db.close()

    def test_agrupa_expedientes_para_varios_lotes(self):
        resultado = expedientes_de_lotes(self.db, [self.lote_uno, self.lote_dos])

        self.assertEqual(resultado[self.lote_uno.id], ["40001", "40002"])
        self.assertEqual(resultado[self.lote_dos.id], ["40002"])


class LotesParcialesTest(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()
        expediente = Expediente(numero="57725")
        lote = LoteCarga(
            anio="2026", ejecutora="785", pecosas_solicitadas="4601,4602",
        )
        self.db.add_all([expediente, lote])
        self.db.flush()
        self.pecosa_procesada = Pecosa(
            numero="4601", expediente_id=expediente.id, estado="Normalizada",
        )
        self.pecosa_faltante = Pecosa(
            numero="4602", expediente_id=expediente.id, estado="Recibida",
        )
        self.db.add_all([self.pecosa_procesada, self.pecosa_faltante])
        self.db.flush()
        self.db.add(BienAlta(
            pecosa_id=self.pecosa_procesada.id,
            lote_id=lote.id,
            codigo_patrimonial="740899502020",
            descripcion="BIEN PROCESADO",
        ))
        self.db.commit()
        self.lote = lote

    def tearDown(self):
        self.db.close()

    def test_difiere_pecosa_sin_filas_y_completa_el_lote_parcial(self):
        bienes = self.db.query(BienAlta).filter(BienAlta.lote_id == self.lote.id).all()
        self.assertEqual(_pecosas_no_encontradas(self.lote, bienes), ["4602"])

        correcto, _ = _diferir_pecosa_faltante(
            self.db, self.lote, self.pecosa_faltante.numero,
        )

        self.assertTrue(correcto)
        self.assertEqual(self.lote.pecosas_solicitadas, "4601")
        self.assertEqual(self.pecosa_faltante.estado, "Recibida")
        self.assertEqual(_pecosas_no_encontradas(self.lote, bienes), [])

    def test_no_permite_dejar_un_lote_sin_pecosas_procesadas(self):
        self.lote.pecosas_solicitadas = "4602"
        self.db.commit()

        correcto, mensaje = _diferir_pecosa_faltante(
            self.db, self.lote, self.pecosa_faltante.numero,
        )

        self.assertFalse(correcto)
        self.assertIn("sin ninguna pecosa procesada", mensaje)
        self.assertEqual(self.lote.pecosas_solicitadas, "4602")


class CompletarLoteParcialTest(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()

        self.expediente = Expediente(numero="60150")
        self.persona = Persona(nombre_completo="RESPONSABLE HUAGAL", dni="12345678")
        self.centro = CentroCosto(nombre_depend="HUAGAL", ipress="4501")
        self.lote = LoteCarga(
            anio="2026",
            ejecutora="785",
            pecosas_solicitadas="4744",
            archivo_generado="formato_importacion_lote_26.xls",
        )
        self.db.add_all([self.expediente, self.persona, self.centro, self.lote])
        self.db.flush()
        self.pecosa = Pecosa(
            numero="4744", expediente_id=self.expediente.id, estado="Normalizada",
        )
        self.db.add(self.pecosa)
        self.db.flush()
        self.db.add_all([
            BienAlta(
                pecosa_id=self.pecosa.id, lote_id=self.lote.id,
                codigo_patrimonial="740899502001", descripcion="BIEN EXISTENTE 1",
                persona_id=self.persona.id, centro_costo_id=self.centro.id,
            ),
            BienAlta(
                pecosa_id=self.pecosa.id, lote_id=self.lote.id,
                codigo_patrimonial="740899502002", descripcion="BIEN EXISTENTE 2",
                persona_id=self.persona.id, centro_costo_id=self.centro.id,
            ),
        ])
        self.db.commit()

    def tearDown(self):
        self.db.close()

    def _reporte_corregido(self):
        filas = [
            ("740899502001", "BIEN CAMBIADO", "4744"),
            ("740899502002", "BIEN EXISTENTE 2", "PECOSA 4744-2026"),
            ("740899502003", "BIEN NUEVO", "4744"),
            ("740899502003", "BIEN NUEVO REPETIDO", "4744"),
            ("740899509999", "OTRA PECOSA", "9999"),
        ]
        return pd.DataFrame([
            {
                "codigo_patrimonial": codigo,
                "descripcion": descripcion,
                "observaciones": observacion,
                "modelo": "MODELO",
                "marca": "MARCA",
                "estado_conserv": "1",
                "nro_serie": "SERIE",
                "fecha_movimto": "2026-08-31",
                "nombre_depend": "HUAGAL",
                "nombre_completo": "RESPONSABLE HUAGAL",
            }
            for codigo, descripcion, observacion in filas
        ])

    def test_agrega_solo_el_bien_faltante_al_mismo_lote_y_expediente(self):
        resultado = _completar_bienes_lote(
            self.db, self.lote, self._reporte_corregido()
        )
        self.db.commit()

        bienes = self.db.query(BienAlta).order_by(BienAlta.codigo_patrimonial).all()
        self.assertEqual(resultado["agregados"], 1)
        self.assertEqual(resultado["agregados_por_pecosa"], {"4744": 1})
        self.assertEqual(resultado["duplicados_omitidos"], 3)
        self.assertEqual(len(bienes), 3)
        self.assertEqual(bienes[-1].codigo_patrimonial, "740899502003")
        self.assertTrue(all(bien.lote_id == self.lote.id for bien in bienes))
        self.assertTrue(all(bien.pecosa_id == self.pecosa.id for bien in bienes))
        self.assertEqual(self.pecosa.expediente_id, self.expediente.id)
        self.assertEqual(self.db.query(LoteCarga).count(), 1)
        self.assertEqual(self.db.query(Expediente).count(), 1)
        self.assertIsNone(self.lote.archivo_generado)
        self.assertEqual(bienes[0].descripcion, "BIEN EXISTENTE 1")

    def test_repetir_el_reporte_no_duplica_bienes(self):
        _completar_bienes_lote(self.db, self.lote, self._reporte_corregido())
        self.db.commit()

        resultado = _completar_bienes_lote(
            self.db, self.lote, self._reporte_corregido()
        )
        self.db.commit()

        self.assertEqual(resultado["agregados"], 0)
        self.assertEqual(self.db.query(BienAlta).count(), 3)


class RegularizacionCargaInicialTest(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()

        self.persona = Persona(nombre_completo="RESPONSABLE SIGA", dni="12345678")
        self.centro = CentroCosto(nombre_depend="CENTRO SIGA", ipress="4567")
        expediente = Expediente(numero="41724")
        lote_historico = LoteCarga(anio="", ejecutora="", pecosas_solicitadas="2011")
        lote_normal = LoteCarga(anio="2026", ejecutora="785", pecosas_solicitadas="2473")
        self.db.add_all([self.persona, self.centro, expediente, lote_historico, lote_normal])
        self.db.flush()
        pecosa = Pecosa(numero="2011", expediente_id=expediente.id, estado="StickerGenerado")
        self.db.add(pecosa)
        self.db.flush()
        self.bien_historico = BienAlta(
            pecosa_id=pecosa.id, lote_id=lote_historico.id,
            codigo_patrimonial="602287628807", descripcion="BIEN HISTORICO",
        )
        self.bien_no_encontrado = BienAlta(
            pecosa_id=pecosa.id, lote_id=lote_historico.id,
            codigo_patrimonial="602287628808", descripcion="BIEN SIN FILA",
        )
        self.bien_normal = BienAlta(
            pecosa_id=pecosa.id, lote_id=lote_normal.id,
            codigo_patrimonial="602287628809", descripcion="BIEN NORMAL",
        )
        self.db.add_all([self.bien_historico, self.bien_no_encontrado, self.bien_normal])
        self.db.commit()

    def tearDown(self):
        self.db.close()

    def test_regulariza_todos_los_lotes_historicos_por_codigo_patrimonial(self):
        reporte = pd.DataFrame([{
            "codigo_patrimonial": "602287628807",
            "ano_eje": "2026",
            "sec_ejec": "785",
            "nombre_completo": "RESPONSABLE SIGA",
            "nombre_depend": "CENTRO SIGA",
            "fecha_movimto": "2026-09-02",
            "estado_conserv": "1",
        }])

        resumen = _regularizar_bienes_historicos(self.db, reporte)
        self.db.commit()

        self.assertEqual(resumen["pendientes"], 2)
        self.assertEqual(resumen["actualizados"], 1)
        self.assertEqual(resumen["no_encontrados"], 1)
        self.assertEqual(self.bien_historico.nombre_completo_siga, "RESPONSABLE SIGA")
        self.assertEqual(self.bien_historico.nombre_depend_siga, "CENTRO SIGA")
        self.assertEqual(self.bien_historico.persona_id, self.persona.id)
        self.assertEqual(self.bien_historico.centro_costo_id, self.centro.id)
        self.assertEqual(self.bien_historico.lote.anio, "2026")
        self.assertEqual(self.bien_historico.lote.ejecutora, "785")
        self.assertEqual(resumen["lotes_actualizados"], 1)
        self.assertIsNone(self.bien_normal.nombre_completo_siga)
        self.assertEqual(_resumen_lote(self.db, self.bien_historico.lote)["estado"], "Incompleto")


class ControlImpresionInventarioTest(unittest.TestCase):
    ENCABEZADOS = [
        "Código Patrimonial", "Código QR", "Ruta QR", "Bien",
        "Establecimiento", "RED", "Área", "Marca", "Modelo", "Color",
        "Nr. Serie",
    ]

    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()
        expediente = Expediente(numero="90001")
        lote = LoteCarga(anio="2026", ejecutora="785")
        self.db.add_all([expediente, lote])
        self.db.flush()
        pecosa = Pecosa(numero="9901", expediente_id=expediente.id)
        self.db.add(pecosa)
        self.db.flush()
        self.db.add(BienAlta(
            pecosa_id=pecosa.id,
            lote_id=lote.id,
            codigo_patrimonial="0001",
            descripcion="BIEN CON ALTA",
        ))
        self.db.commit()
        self.rutas = []

    def tearDown(self):
        self.db.close()
        for ruta in self.rutas:
            if os.path.exists(ruta):
                os.remove(ruta)

    def _reporte(self, filas):
        libro = OpenpyxlWorkbook()
        hoja = libro.active
        hoja.append(self.ENCABEZADOS)
        for fila in filas:
            hoja.append(fila)
        temporal = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        temporal.close()
        libro.save(temporal.name)
        self.rutas.append(temporal.name)
        return temporal.name

    def test_importa_actualiza_y_conserva_sin_area(self):
        ruta = self._reporte([
            [
                " 0001 ", 1, "https://sir.example/qr/equipo/1", "EQUIPO 1",
                "ESTABLECIMIENTO A", None, None, "MARCA", "MODELO", "NEGRO", "S1",
            ],
            [
                "0002", 2, None, "EQUIPO 2", "ESTABLECIMIENTO A",
                "RED A", "AREA A", None, None, None, None,
            ],
            [
                "0001", 1, "https://sir.example/qr/equipo/1", "DUPLICADO",
                "ESTABLECIMIENTO A", None, None, None, None, None, None,
            ],
        ])

        resultado = importar_reporte_inventario(
            self.db, ruta, "reporte.xlsx", anio="2026"
        )

        self.assertEqual(resultado["total"], 2)
        self.assertEqual(resultado["duplicados"], 1)
        bienes = self.db.query(BienInventarioImpresion).order_by(
            BienInventarioImpresion.codigo_patrimonial
        ).all()
        self.assertEqual(bienes[0].codigo_patrimonial, "0001")
        self.assertEqual(bienes[0].codigo_qr, "1")
        self.assertIsNone(bienes[0].area)
        self.assertEqual(bienes[0].relacion_alta, "Vinculado")
        self.assertEqual(bienes[1].motivo_bloqueo, "Sin Ruta QR")
        self.assertEqual(bienes[0].estado_impresion, "Pendiente")
        self.assertEqual(bienes[1].estado_impresion, "Bloqueado")

        bienes[0].estado_impresion = "Impreso"
        bienes[0].impreso_en = datetime.utcnow()
        self.db.commit()

        ruta_actualizada = self._reporte([[
            "0001", 1, "https://sir.example/qr/equipo/1", "EQUIPO ACTUALIZADO",
            "ESTABLECIMIENTO A", "RED A", "AREA COMPLETA", "MARCA", "MODELO",
            "NEGRO", "S1",
        ]])
        segundo = importar_reporte_inventario(
            self.db, ruta_actualizada, "reporte_actualizado.xlsx", anio="2026"
        )
        self.assertEqual(segundo["nuevos"], 0)
        self.assertEqual(segundo["actualizados"], 1)
        self.assertEqual(segundo["inactivos"], 1)
        self.db.expire_all()
        primero = self.db.query(BienInventarioImpresion).filter_by(
            codigo_patrimonial="0001"
        ).one()
        segundo_bien = self.db.query(BienInventarioImpresion).filter_by(
            codigo_patrimonial="0002"
        ).one()
        self.assertEqual(primero.area, "AREA COMPLETA")
        self.assertEqual(primero.activo, 1)
        self.assertEqual(primero.estado_impresion, "Impreso")
        self.assertEqual(segundo_bien.activo, 0)

    def test_estado_cambia_al_generar_y_confirmar_impresion(self):
        inventario = InventarioImpresion(
            nombre="Inventario 2026", anio="2026", unidad_ejecutora="DIRESA"
        )
        bien = BienInventarioImpresion(
            inventario=inventario,
            codigo_patrimonial="1001",
            codigo_qr="10",
            ruta_qr="https://sir.example/qr/equipo/10",
            descripcion="BIEN",
            imprimible=1,
        )
        self.db.add_all([inventario, bien])
        self.db.commit()
        self.assertEqual(_estado_bien(bien), "Pendiente")

        lote = LoteImpresionInventario(
            inventario_id=inventario.id, estado="Preparado", total_bienes=1
        )
        item = ItemLoteImpresionInventario(lote=lote, bien=bien)
        self.db.add_all([lote, item])
        self.db.commit()
        self.assertEqual(_estado_bien(bien), "Pendiente")

        _marcar_pdf_generado(self.db, lote)
        self.db.refresh(bien)
        self.assertEqual(_estado_bien(bien), "Sticker generado")
        self.assertIsNotNone(bien.sticker_generado_en)

        consulta = self.db.query(ItemLoteImpresionInventario).filter_by(
            lote_id=lote.id
        )
        actualizados = _confirmar_items_impresos(self.db, lote, consulta)
        self.assertEqual(actualizados, 1)
        self.db.refresh(bien)
        self.assertEqual(_estado_bien(bien), "Impreso")
        self.assertIsNotNone(bien.impreso_en)

    def test_opciones_buscables_respetan_dependencias_y_limite(self):
        inventario = InventarioImpresion(
            nombre="Inventario 2026", anio="2026", unidad_ejecutora="DIRESA"
        )
        self.db.add(inventario)
        self.db.flush()
        for numero, (red, establecimiento, area) in enumerate([
            ("RED A", "CENTRO NORTE", "ALMACEN"),
            ("RED A", "CENTRO SUR", "PATRIMONIO"),
            ("RED B", "CENTRO ESTE", "ALMACEN"),
            ("RED A", None, None),
        ], start=1):
            self.db.add(BienInventarioImpresion(
                inventario_id=inventario.id,
                codigo_patrimonial=f"B{numero}",
                descripcion="BIEN",
                red=red,
                establecimiento=establecimiento,
                area=area,
            ))
        self.db.commit()

        opciones = _opciones_distintas(
            self.db,
            BienInventarioImpresion.establecimiento,
            inventario.id,
            q="sur",
            red="RED A",
        )
        self.assertEqual(opciones, [{"value": "CENTRO SUR", "label": "CENTRO SUR"}])

        con_vacio = _opciones_distintas(
            self.db,
            BienInventarioImpresion.establecimiento,
            inventario.id,
            red="RED A",
        )
        self.assertEqual(con_vacio[0], {"value": "__SIN_DATO__", "label": "Sin dato"})


class MaestroPatrimonialTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.factory = sessionmaker(bind=self.engine)
        self.db = self.factory()
        self.rutas = []

    def tearDown(self):
        self.db.close()
        for ruta in self.rutas:
            if os.path.exists(ruta):
                os.remove(ruta)

    def _fila(self, codigo="0001", qr="QR-001", usuario="USUARIO A"):
        fila = [None] * len(ENCABEZADOS_SIGA)
        valores = {
            0: codigo,
            1: "AUTOCLAVE DE PRUEBA",
            3: "DEPENDENCIA DE PRUEBA",
            5: usuario,
            8: 1250.50,
            9: date(2026, 1, 10),
            10: 1250.50,
            13: "LABORATORIO",
            17: qr,
            19: "100",
            21: 900,
            26: "MARCA A",
            28: "Bueno",
        }
        for indice, valor in valores.items():
            fila[indice] = valor
        return fila

    def _reporte(self, filas, encabezados=None):
        libro = OpenpyxlWorkbook()
        hoja = libro.active
        hoja.append(encabezados or ENCABEZADOS_SIGA)
        for fila in filas:
            hoja.append(fila)
        temporal = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        temporal.close()
        libro.save(temporal.name)
        libro.close()
        self.rutas.append(temporal.name)
        return temporal.name

    def test_separa_firmante_y_completa_dni_desde_el_maestro(self):
        self.db.add(Persona(nombre_completo="PERSONA DE PRUEBA", dni="12345678"))
        self.db.commit()

        self.assertEqual(
            _datos_firmante(self.db, "OTRA PERSONA (87654321)"),
            ("OTRA PERSONA", "87654321"),
        )
        self.assertEqual(
            _datos_firmante(self.db, "PERSONA DE PRUEBA"),
            ("PERSONA DE PRUEBA", "12345678"),
        )

    def test_busquedas_principales_se_aplican_por_separado(self):
        carga = CargaPatrimonial(
            nombre_archivo="a.xlsx", huella_archivo="x" * 64,
            usuario_carga="admin", estado="Completada",
        )
        self.db.add(carga)
        self.db.flush()
        for codigo, qr, descripcion in [
            ("740001", "QR-100", "AUTOCLAVE GRANDE"),
            ("740002", "QR-200", "COMPUTADORA"),
        ]:
            self.db.add(BienPatrimonial(
                codigo_patrimonial=codigo, codigo_qr=qr,
                descripcion=descripcion, nombre_dependencia="DEPENDENCIA",
                usuario="USUARIO", valor_compra=1, valor_inicial=1,
                ubicacion_fisica="OFICINA", valor_neto=1,
                marca="MARCA", estado_conservacion="Bueno",
                datos_importados="{}", datos_fuente="[]",
                ultima_carga_id=carga.id,
            ))
        self.db.commit()

        filtros = {
            "codigo_patrimonial": "740001", "codigo_qr": "QR-100",
            "descripcion": "autoclave", "q": "",
        }
        resultado = _aplicar_filtros(
            self.db.query(BienPatrimonial), filtros
        ).all()

        self.assertEqual([bien.codigo_patrimonial for bien in resultado], ["740001"])

        solo_prefijo_qr = _aplicar_filtros(
            self.db.query(BienPatrimonial),
            {
                "codigo_patrimonial": "", "codigo_qr": "QR-1",
                "descripcion": "", "q": "",
            },
        ).all()
        self.assertEqual(solo_prefijo_qr, [])

    def _validar(self, ruta, tipo_carga="Completa"):
        carga = preparar_carga(
            self.db, ruta, "reporte_mp.xlsx", "admin",
            tipo_carga=tipo_carga,
        )
        with patch("app.services.maestro_patrimonial.SessionLocal", self.factory):
            validar_carga_patrimonial(carga.id, ruta)
        self.db.expire_all()
        return self.db.get(CargaPatrimonial, carga.id)

    def test_valida_confirma_y_clasifica_una_carga(self):
        carga = self._validar(self._reporte([
            self._fila("0001", "QR-001"),
            self._fila("0002", "QR-002"),
        ]))

        self.assertEqual(carga.estado, "Lista para confirmar")
        self.assertEqual(carga.total_nuevos, 2)
        with patch("app.services.maestro_patrimonial.SessionLocal", self.factory):
            confirmar_carga_patrimonial(carga.id)
        self.db.expire_all()
        carga = self.db.get(CargaPatrimonial, carga.id)
        self.assertEqual(carga.estado, "Completada")
        self.assertEqual(self.db.query(BienPatrimonial).count(), 2)
        self.assertEqual(self.db.query(VersionBienPatrimonial).count(), 2)

        segunda = self._validar(self._reporte([
            self._fila("0001", "QR-001", usuario="USUARIO B"),
        ]))
        self.assertEqual(segunda.total_actualizados, 1)
        self.assertEqual(segunda.total_no_incluidos, 1)
        with patch("app.services.maestro_patrimonial.SessionLocal", self.factory):
            confirmar_carga_patrimonial(segunda.id)
        self.db.expire_all()
        bien = self.db.query(BienPatrimonial).filter_by(codigo_patrimonial="0001").one()
        self.assertEqual(bien.usuario, "USUARIO B")
        cambio = self.db.query(CambioBienPatrimonial).filter_by(campo="usuario").one()
        self.assertEqual(cambio.valor_anterior, "USUARIO A")
        self.assertEqual(cambio.valor_nuevo, "USUARIO B")
        no_incluido = self.db.query(BienCargaPatrimonial).filter_by(
            carga_id=segunda.id, codigo_patrimonial="0002"
        ).one()
        self.assertEqual(no_incluido.clasificacion, "No incluido")
        self.assertIsNotNone(no_incluido.bien_id)

    def test_carga_parcial_no_clasifica_ausentes_como_no_incluidos(self):
        primera = self._validar(self._reporte([
            self._fila("0001", "QR-001"),
            self._fila("0002", "QR-002"),
        ]))
        with patch("app.services.maestro_patrimonial.SessionLocal", self.factory):
            confirmar_carga_patrimonial(primera.id)

        parcial = self._validar(
            self._reporte([self._fila("0001", "QR-001")]),
            tipo_carga="Parcial",
        )

        self.assertEqual(parcial.tipo_carga, "Parcial")
        self.assertEqual(parcial.total_no_incluidos, 0)
        self.assertEqual(
            self.db.query(BienCargaPatrimonial).filter_by(
                carga_id=parcial.id, clasificacion="No incluido"
            ).count(),
            0,
        )

    def test_validacion_puede_reanudarse_desde_el_archivo_persistido(self):
        ruta = self._reporte([self._fila("0001", "QR-001")])
        carga = preparar_carga(
            self.db, ruta, "reporte_mp.xlsx", "admin", tipo_carga="Completa"
        )
        self.assertTrue(carga.archivo_contenido)
        os.remove(ruta)

        with patch("app.services.maestro_patrimonial.SessionLocal", self.factory):
            validar_carga_patrimonial_desde_bd(carga.id)

        self.db.expire_all()
        carga = self.db.get(CargaPatrimonial, carga.id)
        self.assertEqual(carga.estado, "Lista para confirmar")
        self.assertEqual(carga.progreso, 100)
        self.assertIsNone(carga.archivo_contenido)

    def test_exportacion_configurable_se_genera_y_conserva_en_base(self):
        carga = self._validar(self._reporte([self._fila("0001", "QR-001")]))
        with patch("app.services.maestro_patrimonial.SessionLocal", self.factory):
            confirmar_carga_patrimonial(carga.id)
        exportacion = ExportacionPatrimonial(
            usuario="admin",
            estado="Pendiente",
            progreso=0,
            tipo_reporte="Resumen personalizado",
            columnas=json.dumps([
                "codigo_patrimonial", "descripcion", "nombre_dependencia"
            ]),
            filtros=json.dumps({"descripcion": "autoclave"}),
        )
        self.db.add(exportacion)
        self.db.commit()

        with patch("app.routers.maestro_patrimonial.SessionLocal", self.factory):
            generar_exportacion_patrimonial(exportacion.id)

        self.db.expire_all()
        exportacion = self.db.get(ExportacionPatrimonial, exportacion.id)
        self.assertEqual(exportacion.estado, "Completada")
        self.assertEqual(exportacion.progreso, 100)
        self.assertEqual(exportacion.total_filas, 1)
        libro = load_workbook(io.BytesIO(exportacion.archivo_contenido))
        self.assertEqual(
            [celda.value for celda in libro["Bienes"][1]],
            ["Código patrimonial", "Descripción", "Dependencia"],
        )
        self.assertGreater(libro["Bienes"].column_dimensions["B"].width, 20)
        libro.close()

    def test_panel_calidad_detecta_qr_faltantes_y_repetidos(self):
        carga = self._validar(self._reporte([
            self._fila("0001", "QR-001"),
            self._fila("0002", "QR-001"),
            self._fila("0003", None),
        ]))
        with patch("app.services.maestro_patrimonial.SessionLocal", self.factory):
            confirmar_carga_patrimonial(carga.id)

        resumen = _resumen_calidad_datos(self.db)

        self.assertEqual(resumen["sin_qr_total"], 1)
        self.assertEqual(resumen["duplicados_total"], 1)

    def test_rechaza_obligatorio_vacio_pero_qr_repetido_es_alerta(self):
        fila_vacia = self._fila("0001", "QR-001")
        fila_vacia[3] = None
        carga = self._validar(self._reporte([
            fila_vacia,
            self._fila("0002", "QR-001"),
            self._fila("0003", "QR-001"),
        ]))

        self.assertEqual(carga.estado, "Rechazada")
        self.assertEqual(carga.total_errores, 1)
        self.assertEqual(carga.total_alertas, 1)
        detalle = json.loads(carga.detalle_validacion)
        self.assertTrue(any(e["campo"] == "Dependencia" for e in detalle))
        alertas = json.loads(carga.detalle_alertas)
        self.assertEqual(alertas[0]["qr"], "QR-001")
        self.assertEqual(self.db.query(BienCargaPatrimonial).count(), 0)

    def test_numero_orden_puede_estar_vacio(self):
        fila = self._fila("0001", "QR-001")
        fila[19] = None

        carga = self._validar(self._reporte([fila]))

        self.assertEqual(carga.estado, "Lista para confirmar")
        self.assertEqual(carga.total_errores, 0)
        with patch("app.services.maestro_patrimonial.SessionLocal", self.factory):
            confirmar_carga_patrimonial(carga.id)
        self.db.expire_all()
        bien = self.db.query(BienPatrimonial).one()
        self.assertIsNone(bien.numero_orden)

    def test_fecha_alta_puede_estar_vacia(self):
        fila = self._fila("0001", "QR-001")
        fila[9] = None

        carga = self._validar(self._reporte([fila]))

        self.assertEqual(carga.estado, "Lista para confirmar")
        self.assertEqual(carga.total_errores, 0)

    def test_qr_repetido_permite_confirmar_y_genera_alerta(self):
        carga = self._validar(self._reporte([
            self._fila("0001", "QR-001"),
            self._fila("0002", "QR-001"),
        ]))

        self.assertEqual(carga.estado, "Lista para confirmar")
        self.assertEqual(carga.total_errores, 0)
        self.assertEqual(carga.total_alertas, 1)
        self.assertEqual(carga.total_nuevos, 2)
        alerta = json.loads(carga.detalle_alertas)[0]
        self.assertEqual(alerta["codigos"], ["0001", "0002"])
        with patch("app.services.maestro_patrimonial.SessionLocal", self.factory):
            confirmar_carga_patrimonial(carga.id)
        self.db.expire_all()
        self.assertEqual(
            self.db.query(BienPatrimonial).filter_by(codigo_qr="QR-001").count(),
            2,
        )

    def test_qr_numerico_elimina_ceros_iniciales_y_se_busca_normalizado(self):
        carga = self._validar(self._reporte([
            self._fila("0001", "00000054"),
        ]))
        with patch("app.services.maestro_patrimonial.SessionLocal", self.factory):
            confirmar_carga_patrimonial(carga.id)
        self.db.expire_all()

        bien = self.db.query(BienPatrimonial).one()
        self.assertEqual(bien.codigo_qr, "54")
        resultado = _aplicar_filtros(
            self.db.query(BienPatrimonial),
            {
                "codigo_patrimonial": "", "codigo_qr": "00000054",
                "descripcion": "", "q": "",
            },
        ).all()
        self.assertEqual([item.id for item in resultado], [bien.id])

    def test_genera_ficha_pdf_sin_incluir_la_importacion(self):
        carga = self._validar(self._reporte([
            self._fila("0001", "00000054"),
        ]))
        with patch("app.services.maestro_patrimonial.SessionLocal", self.factory):
            confirmar_carga_patrimonial(carga.id)
        self.db.expire_all()
        bien = self.db.query(BienPatrimonial).one()

        contenido = generar_ficha_activo_pdf(
            bien, [], [], list(bien.versiones),
        )

        self.assertTrue(contenido.startswith(b"%PDF-"))
        self.assertGreater(len(contenido), 3000)

    def test_edicion_exige_motivo_y_guarda_correccion(self):
        datos = {
            "codigo_patrimonial": "0001", "codigo_qr": "QR-001",
            "descripcion": "AUTOCLAVE", "nombre_dependencia": "DEPENDENCIA",
            "usuario": "USUARIO A", "fecha_compra": None,
            "valor_compra": "100", "fecha_alta": "2026-01-10",
            "valor_inicial": "100", "ubicacion_fisica": "LABORATORIO",
            "modelo": None, "numero_orden": "1", "medidas": None,
            "valor_neto": "80", "numero_documento": None, "marca": "MARCA",
            "estado_conservacion": "Bueno", "fecha_nea": None,
            "numero_serie": None, "color": None, "caracteristicas": None,
            "observaciones": None,
        }
        carga = CargaPatrimonial(
            nombre_archivo="a.xlsx", huella_archivo="x" * 64,
            usuario_carga="admin", estado="Completada",
        )
        self.db.add(carga)
        self.db.flush()
        bien = BienPatrimonial(
            codigo_patrimonial="0001", codigo_qr="QR-001",
            descripcion="AUTOCLAVE", nombre_dependencia="DEPENDENCIA",
            usuario="USUARIO A", valor_compra=100,
            fecha_alta=date(2026, 1, 10), valor_inicial=100,
            ubicacion_fisica="LABORATORIO", numero_orden="1", valor_neto=80,
            marca="MARCA", estado_conservacion="Bueno",
            datos_importados=json.dumps(datos), datos_fuente="[]",
            ultima_carga_id=carga.id,
        )
        self.db.add(bien)
        self.db.commit()
        valores = {campo: getattr(bien, campo) for campo in CAMPOS_EDITABLES}
        valores["usuario"] = "USUARIO CORREGIDO"

        with self.assertRaisesRegex(ValueError, "motivo"):
            editar_bien_patrimonial(self.db, bien, valores, "", "admin")
        cantidad = editar_bien_patrimonial(
            self.db, bien, valores, "Corrección verificada", "admin"
        )

        self.assertEqual(cantidad, 1)
        self.assertEqual(bien.usuario, "USUARIO CORREGIDO")
        correccion = self.db.query(CorreccionBienPatrimonial).one()
        self.assertEqual(correccion.motivo, "Corrección verificada")
        self.assertEqual(correccion.activa, 1)

    def test_nueva_carga_no_pisa_correccion_manual_y_crea_conflicto(self):
        primera = self._validar(self._reporte([self._fila()]))
        with patch("app.services.maestro_patrimonial.SessionLocal", self.factory):
            confirmar_carga_patrimonial(primera.id)
        self.db.expire_all()
        bien = self.db.query(BienPatrimonial).one()
        valores = {campo: getattr(bien, campo) for campo in CAMPOS_EDITABLES}
        valores["descripcion"] = "AUTOCLAVE CORREGIDO MANUALMENTE"
        editar_bien_patrimonial(
            self.db, bien, valores, "Corrección documental", "admin"
        )

        fila_nueva = self._fila()
        fila_nueva[1] = "AUTOCLAVE ACTUALIZADO EN SIGA"
        segunda = self._validar(self._reporte([fila_nueva]))
        with patch("app.services.maestro_patrimonial.SessionLocal", self.factory):
            confirmar_carga_patrimonial(segunda.id)
        self.db.expire_all()
        bien = self.db.query(BienPatrimonial).one()
        conflicto = self.db.query(ConflictoBienPatrimonial).one()
        self.assertEqual(bien.descripcion, "AUTOCLAVE CORREGIDO MANUALMENTE")
        self.assertEqual(conflicto.estado, "Pendiente")
        self.assertEqual(conflicto.valor_siga_nuevo, "AUTOCLAVE ACTUALIZADO EN SIGA")

        resolver_conflicto(self.db, conflicto, "siga", "admin")
        self.db.expire_all()
        bien = self.db.query(BienPatrimonial).one()
        self.assertEqual(bien.descripcion, "AUTOCLAVE ACTUALIZADO EN SIGA")
        self.assertEqual(
            self.db.query(CorreccionBienPatrimonial).one().activa, 0
        )


if __name__ == "__main__":
    unittest.main()
